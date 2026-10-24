"""Annie v3 core
Module 1 = intake (upload / email), text extraction, classification, field extraction, routing.
Module 2 = duplicate check, document linking, case status, GR/SES submission.
Python 3.12."""
from __future__ import annotations

# =====================================================================================
# 1. IMPORTS
# =====================================================================================
import hashlib
import imaplib
import io
import json
import os
import re
import shutil
import sqlite3
import threading
import zipfile
from collections import Counter
from datetime import date, datetime, timedelta
from difflib import SequenceMatcher
from email import message_from_bytes
from email.header import decode_header
from pathlib import Path

import pandas as pd
import pdfplumber

try:
    import numpy as np
except ImportError:
    np = None
try:
    from PIL import Image, ImageOps, ImageSequence
except ImportError:
    Image = ImageOps = ImageSequence = None
try:
    import pytesseract
except ImportError:
    pytesseract = None
try:
    from rapidocr import RapidOCR
except ImportError:
    RapidOCR = None
try:
    import pymupdf
except ImportError:
    pymupdf = None
try:
    from docx import Document
except ImportError:
    Document = None
try:
    from pillow_heif import register_heif_opener
except ImportError:
    register_heif_opener = None
try:
    from dotenv import load_dotenv
except ImportError:
    load_dotenv = None

# =====================================================================================
# 2. CONFIGURATION AND CONSTANTS
# =====================================================================================
BASE = Path(__file__).resolve().parent
if load_dotenv:
    load_dotenv(BASE / ".env", override=True)

DB_PATH = BASE / "annie_v4.db"
FILE_DIR = BASE / "storage" / "files"
ARCHIVE_DIR = BASE / "storage" / "archive"   # zipped case packages (Complete Archive)
REFERENCE_FOLDER = BASE / "reference_samples"      # purchase_orders/, supplier_invoices/, reference_labels.json

# Names of our own company; lines containing these are never mistaken for the supplier.
OWN_ENTITIES = ("northwind holdings",)

# --- Module 2 threshold
LINK_THRESHOLD = 0.80     # single link threshold

# --- Document types (Module 2 names). Module 1 also yields RECEIPT / UNKNOWN.
PO, INV, GR = "PO", "INVOICE", "GR"
REF_FIELD = {PO: "po_reference", INV: "invoice_reference", GR: "ses_gr_number"}

# --- Case statuses
# Invoice cases: S_OPEN_PO -> S_GR -> S_DONE -> S_ARCHIVED.  PO positions: S_OPEN_INV / S_PART / S_BILLED.
S_OPEN_INV, S_OPEN_PO = "Open - Pending Invoice", "Open - Pending PO/PR"
S_PART, S_BILLED = "Partly billed", "Fully billed"
S_CALLOFF, S_CALLED = "Open - Call-Off", "Matched to Call-Off"   # call-off PO/PR cases and the invoices closed into them
S_GR, S_DONE, S_ARCHIVED = "Pending GR/SES", "Complete", "Archived"

# --- Module 1 statuses and routes
M1_READY, M1_GAPS = "READY_FOR_MODULE_2", "READY_WITH_GAPS"
M1_REVIEW, M1_UNREADABLE, M1_UNSUPPORTED = "UNCLASSIFIED_ERROR", "UNREADABLE_DOCUMENT", "UNSUPPORTED_DOCUMENT"
R_M2, R_REVIEW = "MODULE_2", "INTAKE_REVIEW"
R_DUP_FILE, R_DUP_CONTENT = "REJECTED_DUPLICATE_FILE", "REJECTED_DUPLICATE_CONTENT"


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp", ".heic", ".heif"}
SUPPORTED_EXTENSIONS = {".pdf", ".csv", ".xlsx", ".txt", ".docx", *IMAGE_EXTENSIONS}
REFERENCE_EXTENSIONS = {".pdf", ".txt", ".docx", ".json", *IMAGE_EXTENSIONS}
MAX_FILE_SIZE = 20 * 1024 * 1024
MIN_USABLE_TEXT_CHARACTERS = 25

# Signature heuristic for scanned/image forms: fractions of page (left, top, right, bottom)
SIGNATURE_REGION = (0.03, 0.78, 0.50, 0.90)   # fallback box when no signature line is found
SIGNATURE_INK_RATIO = 0.002                   # share of dark pixels in the signature band that counts as a mark
SIGNATURE_BAND = 0.085                        # band height above the signature line, as a share of page height

LABELS = {
    "po_reference": "PO reference number", "invoice_reference": "Invoice reference number",
    "cost_centre": "Cost centre", "date": "Date", "vendor_name": "Vendor name",
    "approver_name": "Approver name", "approver_title": "Approver job title",
    "prepared_by": "Prepared by (name)", "prepared_by_title": "Prepared by (job title)",
    "category": "Goods / Services tick", "signature_present": "Signature",
    "lines": "Goods lines (description, qty, price/unit, total) or service description",
    "total": "Total amount", "ses_gr_number": "SES/GR number",
    "service_date_from": "Performance/service start date", "service_date_to": "Performance/service end date",
}
MANDATORY = {   # Module 2 mandatory fields
    PO: ["po_reference", "cost_centre", "date", "vendor_name", "approver_name", "approver_title",
         "prepared_by", "prepared_by_title", "category", "signature_present", "lines"],
    INV: ["invoice_reference", "vendor_name", "date", "lines", "total"],
    GR: ["po_reference", "ses_gr_number", "service_date_from", "service_date_to", "prepared_by", "lines"],
}

# Module 1 gate (the only gate): a document is READY_FOR_MODULE_2 only when every mandatory field is present.
# Invoices do not need a PO reference.
M1_REQUIRED = {PO: MANDATORY[PO], INV: MANDATORY[INV]}

DEFAULT_REFERENCE_CONFIG = {
    "classification_threshold": 0.18,
    "document_type_aliases": {
        PO: ["purchase requisition", "purchase order document", "purchase order", "requisition", "po", "pr"],
        INV: ["supplier invoice", "sales invoice", "tax invoice", "commercial invoice", "travel invoice", "invoice"],
        "RECEIPT": ["receipt", "payment receipt"],
    },
    "field_labels": {
        "document_type": ["Document Type", "Document Category", "Type"],
        "supplier": ["Supplier", "Supplier Name", "Vendor", "Vendor Name", "Issued By", "Seller", "Sold By", "From"],
        "invoice_number": ["Invoice Number", "Invoice No", "Invoice #", "Inv #", "Inv No", "Sales Invoice Number",
                           "Tax Invoice Number", "Bill Number", "Document Number", "Document No", "Re: Invoice", "Number"],
        "purchase_order_number": ["PO Reference No.", "PO Reference No", "PO Reference Number", "PO Number", "PO No",
                                  "PO #", "Purchase Order Number", "Purchase Order No", "Requisition Number",
                                  "PR Number", "PR No", "Order Number", "Document Number"],
        "po_reference": ["PO Reference(s)", "PO References", "PO Refs", "PO Numbers", "PO Reference", "PO Reference No.", "Purchase Order Reference", "PO Number", "PO No", "PO #",
                         "Purchase Order Number", "Customer PO", "Client PO", "Buyer PO", "Your PO Ref", "Your PO"],
        "document_date": ["Document Date", "Invoice Date", "Order Date", "PO Date", "Issued", "Date"],
        "amount": ["Total Amount", "Invoice Total", "Order Total", "Grand Total", "Amount Due", "Amount Payable",
                   "Total Due", "Total Fee", "Total Charges", "Net Total", "Amount", "Total"],
        "currency": ["Currency", "Currency Code"],
        "cost_centre": ["Cost Centre", "Cost Center", "Cost Ctr"],
        "approver_name": ["Approver Name", "Approved By", "Approver"],
        "approver_title": ["Approver Job Title", "Approver Title", "Approver Designation"],
        "prepared_by": ["Prepared By", "Requested By", "Requestor"],
        "prepared_by_title": ["Prepared By Job Title", "Job Title", "Designation"],
    },
}
# Reference-sample JSON may use Module 2 names for label keys; map them to the config keys.
LABEL_KEY_ALIASES = {"vendor_name": "supplier", "invoice_reference": "invoice_number", "date": "document_date",
                     "total": "amount", "po_reference_number": "purchase_order_number"}
REFERENCE_STOP_WORDS = {"THE", "AND", "FOR", "WITH", "FROM", "THIS", "THAT", "YOUR", "OUR", "PAGE", "PLEASE",
                        "TOTAL", "DATE", "DOCUMENT", "AMOUNT", "NUMBER"}
_RAPID_OCR_ENGINE = None
_REFERENCE_PROFILE_CACHE = {"signature": None, "profiles": []}

# =====================================================================================
# 3. DATABASE
# =====================================================================================
SCHEMA = """
CREATE TABLE IF NOT EXISTS intake(              -- Module 1 register: every record that entered Annie
  record_id TEXT PRIMARY KEY, channel TEXT, source_file TEXT, source_hash TEXT, content_fingerprint TEXT,
  email_id TEXT, file_path TEXT, extraction_method TEXT, ocr_used INTEGER, classification_method TEXT,
  matched_reference_sample TEXT, doc_type TEXT, m1_status TEXT, route TEXT, extracted TEXT, missing TEXT,
  note TEXT, received_at TEXT, resolved_at TEXT);
CREATE TABLE IF NOT EXISTS attachments(source_hash TEXT PRIMARY KEY, stored_path TEXT, first_email_id TEXT,
  original_filename TEXT, received_at TEXT);
CREATE TABLE IF NOT EXISTS content_registry(content_fingerprint TEXT PRIMARY KEY, source_hash TEXT,
  source_file TEXT, email_id TEXT, registered_at TEXT);
CREATE TABLE IF NOT EXISTS documents(           -- Module 2 documents (PO, INVOICE, GR)
  record_id TEXT PRIMARY KEY, doc_type TEXT, payload TEXT, file_path TEXT, received_at TEXT,
  outcome TEXT, case_id TEXT, link_score REAL, duplicate_of TEXT, missing TEXT);
CREATE TABLE IF NOT EXISTS po_lines(            -- one row per PO/PR line with running balances
  line_id INTEGER PRIMARY KEY AUTOINCREMENT, po_record_id TEXT, po_reference TEXT, line_no INTEGER, description TEXT,
  basis TEXT, ordered_qty REAL, unit_price REAL, ordered_amount REAL,
  received_qty REAL DEFAULT 0, received_amount REAL DEFAULT 0, billed_qty REAL DEFAULT 0, billed_amount REAL DEFAULT 0);
CREATE TABLE IF NOT EXISTS allocations(         -- invoice line -> PO line (many-to-many)
  alloc_id INTEGER PRIMARY KEY AUTOINCREMENT, doc_record_id TEXT, case_id TEXT, kind TEXT, po_line_id INTEGER,
  qty REAL, amount REAL, score REAL, created_at TEXT, manual INTEGER DEFAULT 0, reason TEXT);
CREATE TABLE IF NOT EXISTS cases(               -- one case per invoice
  case_id TEXT PRIMARY KEY, inv_record_id TEXT, status TEXT, reason TEXT, link_score REAL, created_at TEXT, updated_at TEXT,
  hold INTEGER DEFAULT 0,   -- hold=1: user unmatched it, so automatic retries leave it alone
  calloff_id TEXT);          -- set when the invoice was closed into a call-off case
CREATE TABLE IF NOT EXISTS calloffs(               -- PO/PRs with no amount cited: manual matching, many invoices
  case_id TEXT PRIMARY KEY, po_record_id TEXT, status TEXT, reason TEXT, created_at TEXT, updated_at TEXT);
CREATE TABLE IF NOT EXISTS calloff_invoices(     -- invoices closed into a call-off case
  id INTEGER PRIMARY KEY AUTOINCREMENT, case_id TEXT, inv_case_id TEXT, inv_record_id TEXT, added_at TEXT);
CREATE TABLE IF NOT EXISTS archive(case_id TEXT PRIMARY KEY, posted_at TEXT, package_path TEXT);
CREATE TABLE IF NOT EXISTS log(ts TEXT, stage TEXT, record_id TEXT, case_id TEXT, event TEXT, detail TEXT);
CREATE TABLE IF NOT EXISTS mail_messages(       -- inbox monitor: one row per IMAP message already handled
  mailbox TEXT NOT NULL, uid_validity TEXT NOT NULL, uid TEXT NOT NULL, status TEXT NOT NULL, subject TEXT,
  received_at TEXT NOT NULL, completed_at TEXT, attachment_count INTEGER NOT NULL DEFAULT 0, detail TEXT,
  PRIMARY KEY(mailbox, uid_validity, uid));
CREATE TABLE IF NOT EXISTS app_alerts(          -- inbox monitor: notices for the user
  id INTEGER PRIMARY KEY AUTOINCREMENT, created_at TEXT NOT NULL, kind TEXT NOT NULL, title TEXT NOT NULL,
  detail TEXT NOT NULL, record_id TEXT, case_id TEXT, acknowledged_at TEXT);
CREATE TABLE IF NOT EXISTS app_state(key TEXT PRIMARY KEY, value TEXT NOT NULL);   -- e.g. last successful inbox scan
"""
TABLES = ["intake", "attachments", "content_registry", "documents", "po_lines", "allocations", "cases", "calloffs", "calloff_invoices", "archive", "log",
          "mail_messages", "app_alerts", "app_state"]

def _migrate(con):
    """Add columns introduced after a database was created (so older .db files keep working)."""
    def cols(t):
        return {r[1] for r in con.execute(f"PRAGMA table_info({t})")}
    if "manual" not in cols("allocations"):
        con.execute("ALTER TABLE allocations ADD COLUMN manual INTEGER DEFAULT 0")
    if "reason" not in cols("allocations"):
        con.execute("ALTER TABLE allocations ADD COLUMN reason TEXT")
    if "hold" not in cols("cases"):
        con.execute("ALTER TABLE cases ADD COLUMN hold INTEGER DEFAULT 0")
    if "calloff_id" not in cols("cases"):
        con.execute("ALTER TABLE cases ADD COLUMN calloff_id TEXT")
    con.commit()

def connect(path=DB_PATH) -> sqlite3.Connection:
    FILE_DIR.mkdir(parents=True, exist_ok=True)
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    for sub in ("purchase_orders", "supplier_invoices"):
        (REFERENCE_FOLDER / sub).mkdir(parents=True, exist_ok=True)
    con = sqlite3.connect(path, check_same_thread=False)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA busy_timeout=5000")   # the inbox monitor thread and the UI share this file
    con.execute("PRAGMA journal_mode=WAL")
    con.executescript(SCHEMA)
    _migrate(con)
    return con

def reset_db(con):
    for t in TABLES:
        con.execute(f"DROP TABLE IF EXISTS {t}")
    con.executescript(SCHEMA)
    shutil.rmtree(FILE_DIR, ignore_errors=True)
    shutil.rmtree(ARCHIVE_DIR, ignore_errors=True)
    FILE_DIR.mkdir(parents=True, exist_ok=True)
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)

def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")

def _log(con, record_id, case_id, event, detail="", stage="MODULE 2"):
    con.execute("INSERT INTO log VALUES(?,?,?,?,?,?)", (_now(), stage, record_id, case_id, event, detail))

# =====================================================================================
# 4. MODULE 1 - HELPERS AND EMAIL INTAKE
# =====================================================================================
def decode_email_text(value):
    if not value:
        return ""
    out = []
    for part, enc in decode_header(value):
        out.append(part.decode(enc or "utf-8", errors="replace") if isinstance(part, bytes) else part)
    return "".join(out)

def safe_filename(filename):
    return re.sub(r"[^A-Za-z0-9._-]", "_", Path(filename).name)

def calculate_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()

def create_record_id(source_hash, source_row=None, number=None):
    return hashlib.sha256(f"{source_hash}|{source_row or 1}|{number or ''}".encode()).hexdigest()[:16]

def calculate_content_fingerprint(text):
    canonical = re.sub(r"[^A-Z0-9]+", "", str(text).upper())
    return hashlib.sha256(canonical.encode()).hexdigest() if len(canonical) >= MIN_USABLE_TEXT_CHARACTERS else None

def email_config_status() -> dict:
    need = ["EMAIL_HOST", "EMAIL_ADDRESS", "EMAIL_PASSWORD"]
    missing = [k for k in need if not os.getenv(k)]
    return dict(ok=not missing, missing=missing, subject_filter=os.getenv("EMAIL_SUBJECT_FILTER", "").strip())

# =====================================================================================
# 5. MODULE 1 - TEXT EXTRACTION (PDF, OCR, DOCX, TXT)
# =====================================================================================
def is_usable_text(text):
    return len(re.sub(r"\W", "", text or "")) >= MIN_USABLE_TEXT_CHARACTERS

def tesseract_available():
    if pytesseract is None:
        return False
    try:
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False

def ocr_engine_name():
    if Image is None:
        return None
    if RapidOCR is not None:
        return "RAPIDOCR_ONNX"
    return "TESSERACT" if tesseract_available() else None

def ocr_available():
    return ocr_engine_name() is not None

def prepare_image_for_ocr(image):
    gray = ImageOps.autocontrast(ImageOps.grayscale(image.convert("RGB")))
    if gray.width < 1800:
        s = 1800 / gray.width
        gray = gray.resize((int(gray.width * s), int(gray.height * s)))
    return gray

def ocr_image(image):
    global _RAPID_OCR_ENGINE
    if not ocr_available():
        raise RuntimeError("OCR dependencies are unavailable. Install the project requirements.")
    prepared = prepare_image_for_ocr(image)
    if RapidOCR is not None:
        if _RAPID_OCR_ENGINE is None:
            _RAPID_OCR_ENGINE = RapidOCR()
        result = _RAPID_OCR_ENGINE(np.asarray(prepared.convert("RGB")))
        lines = getattr(result, "txts", None) or ()
        return "\n".join(str(x) for x in lines if str(x).strip()).strip()
    return pytesseract.image_to_string(prepared, config="--psm 6").strip()

def extract_pdf_text(path):
    with pdfplumber.open(path) as pdf:
        direct = "\n".join((p.extract_text() or "") for p in pdf.pages).strip()
    if is_usable_text(direct):
        return direct, "PDF_TEXT", False
    if pymupdf is None or not ocr_available():
        return direct, "PDF_TEXT_NO_OCR", False
    pages = []
    with pymupdf.open(path) as doc:
        for page in doc:
            pix = page.get_pixmap(matrix=pymupdf.Matrix(2.2, 2.2), alpha=False)
            pages.append(ocr_image(Image.frombytes("RGB", (pix.width, pix.height), pix.samples)))
    return "\n".join(pages).strip(), "PDF_OCR", True

def extract_image_text(path):
    if Image is None or not ocr_available():
        return "", "IMAGE_NO_OCR", False
    if path.suffix.lower() in {".heic", ".heif"} and register_heif_opener:
        register_heif_opener()
    pages = []
    with Image.open(path) as img:
        for frame in ImageSequence.Iterator(img):
            pages.append(ocr_image(frame.copy()))
    return "\n".join(pages).strip(), "IMAGE_OCR", True

def extract_docx_text(path):
    if Document is None:
        raise RuntimeError("python-docx is required to read DOCX files.")
    doc = Document(path)
    out = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            out.append(" ".join(c.text.strip() for c in row.cells))
    return "\n".join(out).strip(), "DOCX_TEXT", False

def extract_text_from_path(path: Path):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf_text(path)
    if ext in IMAGE_EXTENSIONS:
        return extract_image_text(path)
    if ext == ".docx":
        return extract_docx_text(path)
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="replace").strip(), "TEXT_FILE", False
    raise ValueError(f"Text extraction is not supported for {ext}.")

# =====================================================================================
# 6. MODULE 1 - REFERENCE SAMPLES AND CLASSIFICATION
# =====================================================================================
def merge_config(base, override):
    out = dict(base)
    for k, v in override.items():
        out[k] = merge_config(out[k], v) if isinstance(v, dict) and isinstance(out.get(k), dict) else v
    return out

def normalise_reference_document_type(value):
    if value is None:
        return None
    n = re.sub(r"[^A-Z0-9]+", "_", str(value).upper()).strip("_")
    if n in {"PO", "PR", "PURCHASE_ORDER", "PURCHASE_ORDER_DOCUMENT", "PURCHASE_REQUISITION"}:
        return PO
    if n in {"INVOICE", "SALES_INVOICE", "SUPPLIER_INVOICE", "TAX_INVOICE", "COMMERCIAL_INVOICE"}:
        return INV
    return n

def read_json_reference(path, expected=None):
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return None
    if not isinstance(data, dict):
        return None
    dtype = normalise_reference_document_type(data.get("document_type"))
    if expected and dtype and dtype != expected:
        return None
    values, labels, kw = data.get("field_values", {}), data.get("labels", {}), data.get("keywords", [])
    if not isinstance(values, dict) or not isinstance(labels, dict):
        return None
    kw = [kw] if isinstance(kw, str) else kw
    if not isinstance(kw, list):
        return None
    clean = {}
    for field, aliases in labels.items():
        aliases = [aliases] if isinstance(aliases, str) else aliases
        if not isinstance(aliases, list):
            return None
        clean[LABEL_KEY_ALIASES.get(field, field)] = [str(a).strip() for a in aliases if str(a).strip()]
    return dict(document_type=dtype or expected, field_values=values, labels=clean,
                keywords=[str(k).strip() for k in kw if str(k).strip()])

def json_reference_text(ref):
    out = [f"Document Type {ref.get('document_type')}"] if ref.get("document_type") else []
    for f, v in sorted(ref.get("field_values", {}).items()):
        if v is not None and str(v).strip():
            out.append(f"{f.replace('_', ' ')} {v}")
    for f, aliases in sorted(ref.get("labels", {}).items()):
        out += [f"{f.replace('_', ' ')} {a}" for a in aliases]
    return "\n".join(out + ref.get("keywords", []))

def iter_reference_files():
    for dtype, folder in ((PO, REFERENCE_FOLDER / "purchase_orders"), (INV, REFERENCE_FOLDER / "supplier_invoices")):
        if folder.exists():
            for p in sorted(folder.iterdir()):
                if p.is_file() and p.suffix.lower() in REFERENCE_EXTENSIONS:
                    yield dtype, p

def load_reference_config():
    config = json.loads(json.dumps(DEFAULT_REFERENCE_CONFIG))
    path = REFERENCE_FOLDER / "reference_labels.json"
    if path.exists():
        try:
            config = merge_config(config, json.loads(path.read_text(encoding="utf-8")))
        except (json.JSONDecodeError, OSError):
            pass
    allowed = config["field_labels"]
    for dtype, p in iter_reference_files():
        if p.suffix.lower() != ".json":
            continue
        ref = read_json_reference(p, dtype)
        for field, aliases in (ref or {}).get("labels", {}).items():
            if field in allowed:
                have = {str(x).casefold() for x in allowed[field]}
                allowed[field] += [a for a in aliases if a.casefold() not in have]
    return config

def reference_tokens(text):
    return set(Counter(t for t in re.findall(r"[A-Z]{3,}", str(text).upper()) if t not in REFERENCE_STOP_WORDS))

def build_reference_profiles():
    files = list(iter_reference_files())
    sig = tuple((str(p), p.stat().st_mtime_ns, p.stat().st_size) for _, p in files)
    if _REFERENCE_PROFILE_CACHE["signature"] == sig:
        return _REFERENCE_PROFILE_CACHE["profiles"]
    profiles = []
    for dtype, p in files:
        try:
            if p.suffix.lower() == ".json":
                ref = read_json_reference(p, dtype)
                if not ref:
                    continue
                text = json_reference_text(ref)
            else:
                text, _, _ = extract_text_from_path(p)
        except Exception:
            continue
        tokens = reference_tokens(text)
        if tokens:
            profiles.append(dict(document_type=dtype, path=p, tokens=tokens))
    _REFERENCE_PROFILE_CACHE.update(signature=sig, profiles=profiles)
    return profiles

def get_reference_summary() -> dict:
    files = list(iter_reference_files())
    return dict(purchase_orders=sum(1 for t, _ in files if t == PO), supplier_invoices=sum(1 for t, _ in files if t == INV),
                ocr_engine=ocr_engine_name(), ocr_available=ocr_available())

_JUNK_WORDS = re.compile(r"\b(total|sgd|usd|eur|gbp|qty|date|amount|rate)\b", re.I)
_DATE_RE = re.compile(r"\b\d{4}-\d{2}-\d{2}\b|\b\d{1,2}[/-]\d{1,2}[/-]\d{4}\b|"
                      r"\b\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}\b", re.I)

def extract_labeled_value(labels, text, stop=None, kind=None, avoid=None):
    """Find the value printed after a label. kind='ref' accepts only reference-like values (must hold a digit and no
    column words such as Total/SGD); kind='date' accepts only dates. If a label sits alone on its line (two-column or
    sidebar layouts), the value is looked for on the next two lines."""
    lines = [l.strip() for l in str(text).splitlines() if l.strip()]
    labels = sorted({str(l).strip() for l in labels if str(l).strip()}, key=len, reverse=True)

    def finish(value):
        value = (value or "").strip()
        if stop is not None:
            m = stop.search(value)
            if m and m.start() > 0:
                value = value[:m.start()].strip()
        if kind == "date":
            m = _DATE_RE.search(value)
            return m.group(0) if m else None
        if kind == "ref":
            if not value or _JUNK_WORDS.search(value) or not re.search(r"\d", value):
                return None
            toks = value.split()
            if len(toks) > 4 and not re.search(r"[,;]", value):
                value = next((t for t in toks if re.search(r"\d", t)), toks[0])
        return value or None

    for label in labels:
        strict = re.compile(rf"^{re.escape(label)}\s*(?:[:#-]\s*|\s{{2,}})(.+)$", re.IGNORECASE)
        for line in lines:
            m = strict.match(line)
            if m and finish(m.group(1)):
                return finish(m.group(1))
        for line in lines:
            if line.startswith(label + " ") and finish(line[len(label):]):
                return finish(line[len(label):])
        for line in lines:   # OCR sometimes drops the space after a label
            if line.startswith(label) and len(line) > len(label):
                nxt = line[len(label)]
                if nxt.isalnum() and not nxt.islower() and finish(line[len(label):]):
                    return finish(line[len(label):])
        for i, line in enumerate(lines):   # label alone on its line: value is on one of the next two lines
            if not line.lower().startswith(label.lower()):
                continue
            rest = line[len(label):]
            alone = not rest.strip(" .:#-")
            if kind is None and not alone:
                continue
            if not alone and rest[:1].isalnum():
                continue
            for nxt_line in lines[i + 1:i + 3]:
                if avoid is not None and avoid.match(nxt_line):
                    break   # the next line is another label (this field is simply blank), not its value
                v = finish(nxt_line)
                if v:
                    return v
    return None

def normalise_type_from_value(value, aliases):
    if not value:
        return "UNKNOWN"
    n = " " + re.sub(r"[^A-Z0-9]+", " ", str(value).upper()).strip() + " "
    for dtype, values in aliases.items():
        for alias in values:
            a = " " + re.sub(r"[^A-Z0-9]+", " ", alias.upper()).strip() + " "
            if a in n:
                return dtype
    return "UNKNOWN"

def classify_document(text, config, profiles):
    labelled = extract_labeled_value(config["field_labels"]["document_type"], text)
    dtype = normalise_type_from_value(labelled, config["document_type_aliases"])
    if dtype != "UNKNOWN":
        return dtype, "EXPLICIT_DOCUMENT_TYPE", None
    lines = [l.strip() for l in str(text).splitlines() if l.strip()][:25]
    for line in lines:
        if re.match(r"^PURCHASE\s+(?:ORDER|REQUISITION)(?:\b|\s*[-#])", line, re.I):
            return PO, "TITLE_PATTERN", None
        if re.match(r"^(?:TAX|SALES|SUPPLIER|COMMERCIAL|TRAVEL)?\s*INVOICE(?:\b|\s*[-#])", line, re.I):
            return INV, "TITLE_PATTERN", None
        if re.match(r"^(?:PAYMENT\s+)?RECEIPT(?:\b|\s*[-#])", line, re.I):
            return "RECEIPT", "TITLE_PATTERN", None
    head = " ".join(lines)
    if re.search(r"\binvoice\b", head, re.I) and not re.search(r"purchase\s+(order|requisition)", head, re.I):
        return INV, "TITLE_KEYWORD", None
    incoming = reference_tokens(text)
    groups = config["field_labels"]
    n_fields = sum(extract_labeled_value(groups[f], text) is not None
                   for f in ("supplier", "invoice_number", "purchase_order_number", "po_reference", "document_date", "amount"))
    best, best_score, best_shared = None, 0.0, 0
    for prof in profiles:
        union = incoming | prof["tokens"]
        shared = len(incoming & prof["tokens"])
        score = shared / len(union) if union else 0.0
        if score > best_score:
            best, best_score, best_shared = prof, score, shared
    if best and n_fields >= 2 and best_shared >= 3 and best_score >= float(config.get("classification_threshold", 0.18)):
        return best["document_type"], f"REFERENCE_SAMPLE_SIMILARITY_{best_score:.2f}", best["path"].name
    return "UNKNOWN", "NO_RELIABLE_CLASSIFICATION", None

# =====================================================================================
# 7. MODULE 1 - FIELD EXTRACTION (labels aligned to Module 2 field names)
# =====================================================================================
_AMT = r"[\d,]+\.\d{2}"
HEADER_RE = re.compile(r"(description|particulars|itinerary|programme|item)", re.I)
HEADER_COL_RE = re.compile(r"(qty|amount|price|fee|charge|rate|total|unit)", re.I)
STOP_RE = re.compile(r"^(sub\s*total|total|grand total|net total|amount (?:due|payable)|prepared by|payment terms|approver|bank)\b", re.I)
GOODS_ROW = re.compile(rf"^(?P<d>.*\S)\s+(?P<q>\d+(?:\.\d+)?)\s+(?P<p>{_AMT})\s+(?P<t>{_AMT})$")
SERVICE_ROW = re.compile(rf"^(?P<d>.*?\S)?\s*(?:[-–]\s+[-–]\s+)?(?P<t>{_AMT})$")
COMPANY_RE = re.compile(r"\b(pte\.?\s*ltd|limited|ltd|llp|llc|inc|corp|co\.?|gmbh|sdn\.?\s*bhd|services|supplies|hardware|"
                        r"media|academy|solutions|trading|enterprises|logistics|catering|furniture|consulting)\b", re.I)
SKIP_RE = re.compile(r"\b(bill(?:ed)? to|customer|client|attention|ship to|deliver to|invoice|purchase|requisition)\b|^to\s*:", re.I)

def _num(s):
    return float(str(s).replace(",", ""))

def extract_line_items(text):
    """Return (items, service_description) from the first description-style table in the text."""
    lines = [l.strip() for l in str(text).splitlines() if l.strip()]
    start = next((i for i, l in enumerate(lines) if HEADER_RE.search(l) and HEADER_COL_RE.search(l)
                  and not re.search(r"\d{3,}", l)), None)
    if start is None:
        return [], None
    items, service = [], []
    service_table = bool(re.search(r"service", lines[start], re.I))   # 'Service Description' header: amounts may be blank
    for line in lines[start + 1:]:
        if STOP_RE.match(line):
            break
        m = GOODS_ROW.match(line)
        if m:
            items.append(dict(description=m["d"].strip(), qty=_num(m["q"]), unit_price=_num(m["p"]), total=_num(m["t"])))
            continue
        m = SERVICE_ROW.match(line)
        if m and (m["d"] or "").strip():
            service.append(m["d"].strip())
        elif not m and (_DATE_RE.fullmatch(line) or (len(line.split()) == 1 and re.search(r"\d", line))):
            continue   # stray date / reference from a neighbouring column, not a wrapped description
        elif not m and not service_table and not service and not any(i.get("total") for i in items) and re.search(r"[A-Za-z]{3}", line):
            mq = re.match(r"^(?P<d>.*?[A-Za-z].*?\S)\s+(?P<q>\d+(?:\.\d+)?)$", line)   # goods row that cites no price
            items.append(dict(description=(mq["d"] if mq else line), qty=_num(mq["q"]) if mq else None, unit_price=None, total=None))
        elif not m and items:
            items[-1]["description"] += " " + line
        elif not m and service:
            service[-1] += " " + line
        elif not m and service_table and re.search(r"[A-Za-z]{3}", line):
            service.append(line)   # service line with no amount printed (PO that cites no price)
    return items, (" ".join(service) if service and not items else None)

def detect_category(text, items, service_description):
    for line in str(text).splitlines():
        if re.search(r"\bgoods\b", line, re.I) and re.search(r"\bservices?\b", line, re.I):
            m = re.search(r"(?:\bX|✓|✔|☑|\[x\])\s*(Goods|Services?)\b", line, re.I)
            if m:
                return "Goods" if m.group(1).lower().startswith("good") else "Services"
    if items:
        return "Goods"
    return "Services" if service_description else None

LEGAL_RE = re.compile(r"\b(pte\.?\s*ltd|limited|ltd|llp|llc|inc|corp|gmbh|sdn\.?\s*bhd)\b", re.I)

def guess_supplier(text):
    """Unlabelled supplier: prefer a header line with a legal suffix (Pte Ltd...), else a company-keyword line."""
    own = [o.strip().lower() for o in OWN_ENTITIES if o and o.strip()]   # blank entries would match every line
    lines = [l.strip() for l in str(text).splitlines() if l.strip()][:15]
    for pattern in (LEGAL_RE, COMPANY_RE):
        for line in lines:
            if any(o in line.lower() for o in own) or SKIP_RE.search(line) or not pattern.search(line):
                continue
            name = re.split(r"\s*[|,]\s*|\s{2,}|\s+(?:UEN|Reg\.?\s*No|Tel|Phone)\b", line)[0]
            name = re.sub(r"^[A-Z]\s+(?=[A-Z][a-z])", "", name).strip(" -:")
            if name:
                return name
    return None

def first_date_in_text(text):
    pat = (r"\b\d{4}-\d{2}-\d{2}\b|\b\d{1,2}[/-]\d{1,2}[/-]\d{4}\b|"
           r"\b\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s+\d{4}\b")
    m = re.search(pat, "\n".join(str(text).splitlines()[:40]), re.I)
    return m.group(0) if m else None

def parse_amount_and_currency(value):
    if not value:
        return None, None
    text = str(value).strip()
    cm = re.search(r"\b(SGD|USD|EUR|GBP|MYR|AUD|JPY|CNY)\b", text, re.I)
    currency = cm.group(1).upper() if cm else None
    if currency is None:
        currency = "SGD" if "S$" in text else "USD" if "$" in text else "EUR" if "€" in text else "GBP" if "£" in text else None
    nm = re.search(r"\(?-?\d[\d,]*(?:\.\d{1,2})?\)?", text)
    if not nm:
        return None, currency
    raw = nm.group(0)
    neg = raw.startswith("(") and raw.endswith(")")
    try:
        amount = float(raw.strip("()").replace(",", ""))
        return round(-amount if neg else amount, 2), currency
    except ValueError:
        return None, currency

_TOTAL_ANY = re.compile(r"\b(?:grand\s+total|total\s+(?:due|fee|charges)|net\s+total|amount\s+(?:due|payable)|total)\b\W{0,3}\s*"
                        r"((?:SGD|USD|EUR|GBP|MYR|AUD|JPY|CNY|S\$|\$)?\s*[\d,]+\.\d{2})\b", re.I)

def find_total_anywhere(text):
    """Fallback when a total is printed mid-line (merged columns): last 'Total ... 0.00' occurrence."""
    hits = _TOTAL_ANY.findall(str(text))
    return hits[-1].strip() if hits else None

def _render_pages(path: Path, max_pages=3):
    """Render pages to images so digital PDFs and scans are treated the same way."""
    ext = path.suffix.lower()
    if Image is None:
        return []
    if ext in IMAGE_EXTENSIONS:
        if ext in (".heic", ".heif") and register_heif_opener:
            register_heif_opener()
        with Image.open(path) as img:
            return [img.convert("RGB")]
    if ext != ".pdf":
        return []
    if pymupdf is not None:
        with pymupdf.open(path) as doc:
            out = []
            for page in list(doc)[:max_pages]:
                pix = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
                out.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
            return out
    with pdfplumber.open(path) as pdf:
        return [p.to_image(resolution=150).original.convert("RGB") for p in pdf.pages[:max_pages]]

def _signature_in_image(img) -> bool:
    """Bottom-left signature check: locate the signature line, then test the band above it for a mark (ink density)
    and run OCR on that band. Returns True if either finds something."""
    gray = ImageOps.grayscale(img)
    w, h = gray.size
    arr = np.asarray(gray) if np is not None else None
    dark = (lambda x, y: arr[y, x] < 160) if arr is not None else (lambda x, y: gray.getpixel((x, y)) < 160)
    line_y = x0 = x1 = None
    for y in range(int(h * 0.55), h):
        run = start = best_len = best_start = 0
        for x in range(int(w * 0.55)):
            if dark(x, y):
                if run == 0:
                    start = x
                run += 1
                if run > best_len:
                    best_len, best_start = run, start
            else:
                run = 0
        if best_len > 0.12 * w:
            line_y, x0, x1 = y, best_start, best_start + best_len
    if line_y is None:
        l, t, r, b = SIGNATURE_REGION
        box = (int(l * w), int(t * h), int(r * w), int(b * h))
    else:
        band = int(SIGNATURE_BAND * h)
        box = (max(x0 - 4, 0), max(line_y - band, 0), min(x1 + 4, w), max(line_y - 4, 1))
    crop = gray.crop(box)
    px = list(crop.getdata())
    if px and sum(1 for p in px if p < 160) / len(px) >= SIGNATURE_INK_RATIO:
        return True
    if ocr_available():
        try:
            return len(re.sub(r"\W", "", ocr_image(crop.convert("RGB")))) >= 2
        except Exception:
            return False
    return False

def detect_signature(path: Path) -> bool:
    """Extra OCR/image step for PO/PR forms: is there a signature in the bottom-left of the page?
    Works on digital PDFs too (pages are rendered first). Always returns True or False."""
    try:
        return any(_signature_in_image(img) for img in _render_pages(path))
    except Exception:
        return False

def split_refs(value):
    """'PR-1, PR-2 and PR-3' -> ['PR-1', 'PR-2', 'PR-3'] (slashes are kept: ZH/26/1187 is one reference)."""
    out = []
    for t in re.split(r"[,;&]|\band\b", str(value or ""), flags=re.I):
        t = t.strip(" .:")
        if t and re.search(r"\d", t) and t.lower() not in {x.lower() for x in out}:
            out.append(t)
    return out

def attach_po_citations(con, raw: dict, text: str) -> dict:
    """Invoices: collect every PO/PR cited (header list, per-line prefix, or a known PO number found anywhere in the text)."""
    if raw.get("doc_type") != INV:
        return raw
    known = []
    for r in con.execute("SELECT payload FROM documents WHERE doc_type=?", (PO,)):
        ref = json.loads(r["payload"]).get("po_reference")
        if ref:
            known.append(ref)
    refs = split_refs(raw.get("po_reference"))
    flat = _ref_norm(text)
    for k in known:
        if _ref_norm(k) in flat and _ref_norm(k) not in {_ref_norm(x) for x in refs}:
            refs.append(k)
    pool = {_ref_norm(x): x for x in refs + known}
    for it in raw.get("items") or []:
        first, _, rest = str(it.get("description", "")).partition(" ")
        if _ref_norm(first) in pool and rest:
            it["po_reference"], it["description"] = pool[_ref_norm(first)], rest.strip()
            if _ref_norm(first) not in {_ref_norm(x) for x in refs}:
                refs.append(pool[_ref_norm(first)])
    raw["po_references"] = refs
    raw["po_reference"] = refs[0] if refs else None
    return raw

def extract_fields_from_text(path: Path, text, method, ocr_used, config, profiles):
    dtype, cmethod, matched = classify_document(text, config, profiles)
    L = config["field_labels"]
    others = {l for v in L.values() for l in v if len(l) >= 3}
    others |= {"Bill To", "Billed to", "Customer", "Client", "To", "Contract", "Attention"}
    stop = re.compile(r"\s+(?:%s)\s*:" % "|".join(re.escape(x) for x in sorted(others, key=len, reverse=True)), re.I)
    kinds = {"po_reference": "ref", "purchase_order_number": "ref", "invoice_number": "ref", "document_date": "date"}
    every = {l for v in L.values() for l in v if len(l) >= 3} | {"Signature", "Approver Signature", "Goods", "Services"}
    avoid = re.compile(r"(?:%s)\b" % "|".join(re.escape(x) for x in sorted(every, key=len, reverse=True)), re.I)
    get = lambda key: extract_labeled_value(L[key], text, stop, kinds.get(key), avoid)

    items, service = extract_line_items(text)
    po_ref = get("po_reference")
    if dtype == PO:
        po_ref = po_ref or get("purchase_order_number")
        inv_ref = None
    else:
        inv_ref = get("invoice_number")
    total, currency = parse_amount_and_currency(get("amount") or find_total_anywhere(text))
    priced = [i["total"] for i in items if i.get("total") is not None]
    if total is None and priced:
        total = round(sum(priced), 2)
    explicit_currency = get("currency")
    return dict(
        doc_type=dtype, vendor_name=get("supplier") or guess_supplier(text),
        po_reference=po_ref, invoice_reference=inv_ref, date=get("document_date") or first_date_in_text(text),
        total=total, currency=explicit_currency or currency,
        cost_centre=get("cost_centre") if dtype == PO else None,
        approver_name=get("approver_name") if dtype == PO else None,
        approver_title=get("approver_title") if dtype == PO else None,
        prepared_by=get("prepared_by") if dtype == PO else None,
        prepared_by_title=get("prepared_by_title") if dtype == PO else None,
        category=detect_category(text, items, service) if dtype == PO else None,
        signature_present=detect_signature(path) if dtype == PO else None,
        items=items, service_description=service,
        source_file=path.name, source_extension=path.suffix.lower(), source_row=None,
        text_extracted=is_usable_text(text), extraction_method=method, ocr_used=ocr_used,
        classification_method=cmethod, matched_reference_sample=matched)

TABLE_ALIASES = {
    "type": "document_type", "document type": "document_type", "document_type": "document_type",
    "vendor": "vendor_name", "vendor name": "vendor_name", "vendor_name": "vendor_name", "supplier": "vendor_name",
    "supplier name": "vendor_name", "supplier_name": "vendor_name",
    "document number": "document_number", "document_number": "document_number",
    "invoice number": "invoice_reference", "invoice_number": "invoice_reference", "invoice no": "invoice_reference",
    "invoice_no": "invoice_reference", "invoice_reference": "invoice_reference",
    "po reference": "po_reference", "po_reference": "po_reference", "po number": "po_reference",
    "po_number": "po_reference", "po no": "po_reference", "po_no": "po_reference",
    "document date": "date", "document_date": "date", "invoice date": "date", "invoice_date": "date", "date": "date",
    "total": "total", "total amount": "total", "total_amount": "total", "amount": "total", "currency": "currency",
    "cost centre": "cost_centre", "cost_centre": "cost_centre", "approver name": "approver_name",
    "approver_name": "approver_name", "approver title": "approver_title", "approver_title": "approver_title",
    "prepared by": "prepared_by", "prepared_by": "prepared_by", "prepared by title": "prepared_by_title",
    "prepared_by_title": "prepared_by_title", "category": "category", "service description": "service_description",
    "service_description": "service_description", "signature": "signature_present", "signature_present": "signature_present",
}

def read_table(path: Path):
    df = pd.read_csv(path) if path.suffix.lower() == ".csv" else pd.read_excel(path)
    df = df.rename(columns={c: TABLE_ALIASES.get(str(c).strip().lower(), str(c).strip().lower().replace(" ", "_")) for c in df.columns})
    method = "CSV_TABLE" if path.suffix.lower() == ".csv" else "EXCEL_TABLE"
    records = []
    for i, row in df.iterrows():
        r = {k: (None if pd.isna(v) else v) for k, v in row.to_dict().items()}
        r.update(source_file=path.name, source_extension=path.suffix.lower(), source_row=i + 2, text_extracted=True,
                 extraction_method=method, ocr_used=False, classification_method="TABLE_FIELD", matched_reference_sample=None)
        records.append(r)
    return records, df.fillna("").astype(str).to_csv(index=False)

# =====================================================================================
# 8. MODULE 1 - CLEANING, VALIDATION AND ROUTING DECISION
# =====================================================================================
def _clean_text(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or not str(v).strip():
        return None
    return re.sub(r"\s+", " ", str(v)).strip()

def _clean_ref(v):
    s = _clean_text(v)
    return s.strip(" .,;:") if s else None

def standardise_date(v):
    s = _clean_text(v)
    if not s:
        return None
    m = re.search(r"\d{4}-\d{2}-\d{2}", s)
    if m:
        try:
            return datetime.strptime(m.group(0), "%Y-%m-%d").strftime("%Y-%m-%d")
        except ValueError:
            pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    parsed = pd.to_datetime(s, errors="coerce", dayfirst=True)
    return None if pd.isna(parsed) else parsed.strftime("%Y-%m-%d")

def standardise_amount(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or not str(v).strip():
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return parse_amount_and_currency(v)[0]

def standardise_signature(v):
    if isinstance(v, bool) or v is None:
        return v
    s = str(v).strip().lower()
    return True if s in ("yes", "y", "true", "1", "signed") else False if s in ("no", "n", "false", "0", "unsigned") else None

def normalise_document_type(value):
    return normalise_type_from_value(_clean_text(value), DEFAULT_REFERENCE_CONFIG["document_type_aliases"])

def clean_record(raw: dict, fingerprint=None, source_hash=None) -> dict:
    dtype = raw.get("doc_type") or normalise_document_type(raw.get("document_type"))
    if dtype not in (PO, INV, "RECEIPT"):
        dtype = "UNKNOWN"
    number = _clean_ref(raw.get("document_number"))
    po_ref, inv_ref = _clean_ref(raw.get("po_reference")), _clean_ref(raw.get("invoice_reference"))
    if dtype == PO:
        po_ref, inv_ref = po_ref or number, None
    elif dtype == INV:
        inv_ref = inv_ref or number
    total = standardise_amount(raw.get("total") if raw.get("total") is not None else raw.get("amount"))
    items = raw.get("items") if isinstance(raw.get("items"), list) else []
    source_hash = source_hash or raw.get("source_hash")
    rec = dict(
        record_id=raw.get("record_id") or create_record_id(source_hash, raw.get("source_row"), inv_ref or po_ref),
        doc_type=dtype, vendor_name=_clean_text(raw.get("vendor_name") or raw.get("supplier")),
        po_reference=po_ref, invoice_reference=inv_ref, po_references=[_clean_ref(x) for x in (raw.get("po_references") or []) if _clean_ref(x)],
        date=standardise_date(raw.get("date") or raw.get("document_date")),
        total=total, currency=(_clean_text(raw.get("currency")) or "").upper() or None,
        cost_centre=_clean_text(raw.get("cost_centre")), approver_name=_clean_text(raw.get("approver_name")),
        approver_title=_clean_text(raw.get("approver_title")), prepared_by=_clean_text(raw.get("prepared_by")),
        prepared_by_title=_clean_text(raw.get("prepared_by_title")), category=_clean_text(raw.get("category")),
        signature_present=standardise_signature(raw.get("signature_present")),
        items=items, service_description=_clean_text(raw.get("service_description")),
        source_file=raw.get("source_file"), source_extension=raw.get("source_extension"), source_hash=source_hash,
        content_fingerprint=fingerprint, source_row=raw.get("source_row"), email_id=raw.get("email_id"),
        extraction_method=raw.get("extraction_method"), ocr_used=bool(raw.get("ocr_used", False)),
        classification_method=raw.get("classification_method"), matched_reference_sample=raw.get("matched_reference_sample"))

    missing = missing_fields(rec) if dtype in (PO, INV) else []
    if not raw.get("text_extracted", True):
        status, note = M1_UNREADABLE, "No usable text was extracted. Correct the document or verify OCR, then recheck."
    elif dtype == "UNKNOWN":
        status, note = M1_REVIEW, "The document type could not be determined reliably. A reviewer must classify it."
    elif dtype not in (PO, INV):
        status, note = M1_UNSUPPORTED, f"{dtype} is readable but outside the current PO/PR and supplier-invoice scope."
    elif missing:
        status, note = M1_GAPS, "Required fields are missing: " + ", ".join(LABELS.get(m, m) for m in missing) + ". A reviewer must complete them."
    else:
        status, note = M1_READY, "The document was extracted and standardised successfully."
    if rec["ocr_used"] and status == M1_READY:
        note = "OCR was used. " + note
    rec.update(processing_status=status, missing_fields=missing, intake_note=note)
    return rec

def to_module2_payload(rec: dict) -> dict:
    """Module 1 record -> the JSON that Module 2 expects."""
    keys = ["record_id", "doc_type", "vendor_name", "date", "total", "currency", "items", "service_description"]
    out = {k: rec.get(k) for k in keys}
    if rec["doc_type"] == PO:
        out.update({k: rec.get(k) for k in ("po_reference", "cost_centre", "approver_name", "approver_title",
                                           "prepared_by", "prepared_by_title", "category", "signature_present")})
    else:
        out.update(invoice_reference=rec.get("invoice_reference"), po_reference=rec.get("po_reference"),
                   po_references=rec.get("po_references") or ([rec["po_reference"]] if rec.get("po_reference") else []))
    return out

# =====================================================================================
# 9. MODULE 1 - INGEST ORCHESTRATION (Submit Document and Email share this)
# =====================================================================================
def _save_intake_row(con, rec, channel, file_path, route, resolved=False):
    rid = rec["record_id"]
    if con.execute("SELECT 1 FROM intake WHERE record_id=?", (rid,)).fetchone():
        rid = rec["record_id"] = f"{rid}-{datetime.now().strftime('%H%M%S%f')}"
    con.execute("INSERT INTO intake VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (rid, channel, rec.get("source_file"), rec.get("source_hash"), rec.get("content_fingerprint"),
                 rec.get("email_id"), str(file_path) if file_path else None, rec.get("extraction_method"),
                 int(bool(rec.get("ocr_used"))), rec.get("classification_method"), rec.get("matched_reference_sample"),
                 rec.get("doc_type"), rec.get("processing_status"), route, json.dumps(rec, default=str),
                 json.dumps(rec.get("missing_fields", [])), rec.get("intake_note"), _now(), _now() if resolved else None))
    return rid

def route_record(con, rec, channel, file_path) -> dict:
    """Routing decision: only READY_FOR_MODULE_2 enters Annie; everything else goes to Intake Review."""
    ready = rec["processing_status"] == M1_READY
    rid = _save_intake_row(con, rec, channel, file_path, R_M2 if ready else R_REVIEW)
    _log(con, rid, None, "EXTRACTED", f"{rec['doc_type']} via {rec['extraction_method']}; ocr={rec['ocr_used']}; "
         f"classified by {rec['classification_method']}", "MODULE 1")
    result = dict(source_file=rec.get("source_file"), record_id=rid, m1_status=rec["processing_status"],
                  note=rec["intake_note"], missing=rec["missing_fields"], channel=channel)
    if ready:
        _log(con, rid, None, "ROUTED_TO_MODULE_2", "READY_FOR_MODULE_2", "MODULE 1")
        payload = to_module2_payload(rec)
        payload["record_id"] = rid
        result.update(route=R_M2, m2=process_document(con, payload, file_path=str(file_path)))
    else:
        _log(con, rid, None, "ROUTED_TO_INTAKE_REVIEW", f"{rec['processing_status']}: {rec['intake_note']}", "MODULE 1")
        result.update(route=R_REVIEW)
        con.commit()
    return result

def _dup_result(con, filename, channel, email_id, source_hash, reason, duplicate_of, stored):
    rec = dict(record_id=create_record_id(source_hash + _now(), None, reason), doc_type="UNKNOWN", source_file=filename,
               source_hash=source_hash, email_id=email_id, processing_status=reason, missing_fields=[],
               intake_note=f"Duplicate of {duplicate_of}; not processed.", extraction_method=None, ocr_used=False)
    route = R_DUP_FILE if reason == "DUPLICATE_FILE" else R_DUP_CONTENT
    rid = _save_intake_row(con, rec, channel, stored, route, resolved=True)
    _log(con, rid, None, "REJECTED_" + reason, f"{filename} duplicates {duplicate_of}", "MODULE 1")
    con.commit()
    return dict(source_file=filename, record_id=rid, m1_status=reason, route=route, note=rec["intake_note"],
                missing=[], channel=channel)

def ingest_file(con, filename: str, content: bytes, channel="UPLOAD", email_id=None) -> list[dict]:
    """Thread-safe entry point: uploads and the inbox monitor take turns, so OCR and duplicate checks never race."""
    with _INGEST_LOCK:
        return _ingest_file(con, filename, content, channel, email_id)

def _ingest_file(con, filename: str, content: bytes, channel="UPLOAD", email_id=None) -> list[dict]:
    """Module 1 for one file: hash check -> store -> extract -> content check -> clean -> route."""
    create = safe_filename(filename)
    ext = Path(create).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS or not content or len(content) > MAX_FILE_SIZE:
        _log(con, None, None, "FILE_REJECTED", f"{filename}: unsupported type, empty or over 20 MB", "MODULE 1")
        con.commit()
        return [dict(source_file=filename, record_id=None, m1_status="FILE_REJECTED", route="REJECTED", missing=[],
                     note="Unsupported file type, empty file, or larger than 20 MB.", channel=channel)]
    h = calculate_hash(content)
    seen = con.execute("SELECT * FROM attachments WHERE source_hash=?", (h,)).fetchone()
    _log(con, None, None, "RECEIVED", f"{filename} via {channel}" + (f" (email {email_id})" if email_id else ""), "MODULE 1")
    if seen:
        return [_dup_result(con, filename, channel, email_id, h, "DUPLICATE_FILE", seen["original_filename"], seen["stored_path"])]
    FILE_DIR.mkdir(parents=True, exist_ok=True)
    path = FILE_DIR / f"{h[:10]}_{create}"
    path.write_bytes(content)
    con.execute("INSERT INTO attachments VALUES(?,?,?,?,?)", (h, str(path), email_id, filename, _now()))

    config, profiles = load_reference_config(), build_reference_profiles()
    try:
        if ext in (".csv", ".xlsx"):
            raws, canonical = read_table(path)
        else:
            text, method, ocr = extract_text_from_path(path)
            canonical = text
            raws = [attach_po_citations(con, extract_fields_from_text(path, text, method, ocr, config, profiles), text)]
        fp = calculate_content_fingerprint(canonical)
        prior = con.execute("SELECT * FROM content_registry WHERE content_fingerprint=?", (fp,)).fetchone() if fp else None
        if prior and prior["source_hash"] != h:
            return [_dup_result(con, filename, channel, email_id, h, "DUPLICATE_CONTENT", prior["source_file"], path)]
        records = [clean_record({**r, "email_id": email_id}, fp, h) for r in raws]
        if fp and not prior:
            con.execute("INSERT INTO content_registry VALUES(?,?,?,?,?)", (fp, h, filename, email_id, _now()))
    except Exception as err:
        records = [clean_record(dict(doc_type="UNKNOWN", source_file=path.name, source_extension=ext, email_id=email_id,
                                     text_extracted=False, extraction_method="PROCESSING_ERROR",
                                     classification_method="PROCESSING_ERROR"), None, h)]
        records[0]["processing_status"], records[0]["intake_note"] = M1_REVIEW, f"Processing error: {err}"
    return [route_record(con, rec, channel, path) for rec in records]

# --- Inbox monitor: poll_inbox() = one scan; start_monitor() = scan every POLL_SECONDS in a background thread.
# .env: EMAIL_HOST / EMAIL_ADDRESS / EMAIL_PASSWORD (required); EMAIL_PORT (993), EMAIL_MAILBOX (INBOX),
#       EMAIL_SUBJECT_FILTER, EMAIL_LOOKBACK_DAYS (30, first scan only), POLL_SECONDS (60, clamped to 15-3600).
DEFAULT_POLL_SECONDS = 60
_POLL_LOCK = threading.Lock()      # one scan at a time (timer and "Check inbox now" never overlap)
_INGEST_LOCK = threading.RLock()   # one file ingested at a time
_MONITOR = dict(thread=None, stop=threading.Event(), guard=threading.Lock(), results=[], alerts=[],
                version=0, last=None, last_run=None, last_error=None)

def poll_seconds() -> int:
    try:
        return max(15, min(3600, int(os.getenv("POLL_SECONDS", DEFAULT_POLL_SECONDS))))
    except ValueError:
        return DEFAULT_POLL_SECONDS

def _save_alert(con, kind, title, detail, record_id=""):
    con.execute("INSERT INTO app_alerts(created_at,kind,title,detail,record_id) VALUES(?,?,?,?,?)",
                (_now(), kind, title, detail, record_id))
    con.commit()
    with _MONITOR["guard"]:
        _MONITOR["alerts"].append(dict(kind=kind, title=title, detail=detail))

def report_monitor_error(con, detail: str):
    """Raise one alert for a continuing connection/configuration error (repeats of the same text are ignored)."""
    prior = con.execute("SELECT detail FROM app_alerts WHERE kind='EMAIL_CONNECTION' AND acknowledged_at IS NULL "
                        "ORDER BY id DESC LIMIT 1").fetchone()
    if not prior or prior["detail"] != detail[:300]:
        _save_alert(con, "EMAIL_CONNECTION", "Inbox monitoring needs attention", detail[:300])

def open_alerts(con, limit=50) -> list[dict]:
    return [dict(r) for r in con.execute("SELECT * FROM app_alerts WHERE acknowledged_at IS NULL ORDER BY id DESC LIMIT ?", (limit,))]

def acknowledge_alert(con, alert_id: int):
    con.execute("UPDATE app_alerts SET acknowledged_at=? WHERE id=?", (_now(), alert_id))
    con.commit()

def _record_message(con, mailbox, uid_validity, uid, status, subject, count, detail):
    now = _now()
    con.execute("""INSERT INTO mail_messages VALUES(?,?,?,?,?,?,?,?,?)
                   ON CONFLICT(mailbox,uid_validity,uid) DO UPDATE SET status=excluded.status, subject=excluded.subject,
                   completed_at=excluded.completed_at, attachment_count=excluded.attachment_count, detail=excluded.detail""",
                (mailbox, uid_validity, uid, status, subject, now, now, count, detail))
    con.commit()

def poll_inbox(con=None, blocking=False):
    """Scan the mailbox once. Every message (read or unread) is handled once per IMAP UID; the mailbox is opened
    read-only and mail is never marked as read. A failed message is retried on the next scan (duplicate checks make
    that safe). Returns (results, stats); if another scan is running and blocking=False, returns ([], zero stats)."""
    cfg = email_config_status()
    if not cfg["ok"]:
        raise ValueError("Missing in .env: " + ", ".join(cfg["missing"]))
    stats = dict(checked=0, emails=0, attachments=0, review=0, duplicates=0, skipped=0, skipped_unsupported=0, errors=0)
    results = []
    if not _POLL_LOCK.acquire(blocking=blocking):
        return results, stats
    own, box = con is None, None
    try:
        con = con or connect()
        host, address = os.environ["EMAIL_HOST"], os.environ["EMAIL_ADDRESS"]
        mailbox = os.getenv("EMAIL_MAILBOX", "INBOX")
        mailbox_key = json.dumps([host.casefold(), address.casefold(), mailbox])   # UIDs are per account + mailbox
        scan_key = f"last_successful_scan:{host}:{address}:{mailbox}"
        flt = cfg["subject_filter"].casefold()
        marker = f"Subject filter:{flt}"
        since = datetime.now() - timedelta(days=max(1, min(365, int(os.getenv("EMAIL_LOOKBACK_DAYS", "30")))))
        last = con.execute("SELECT value FROM app_state WHERE key=?", (scan_key,)).fetchone()
        if last:   # resume after downtime; one day of overlap absorbs server/date differences
            since = min(since, datetime.fromisoformat(last["value"]) - timedelta(days=1))
        box = imaplib.IMAP4_SSL(host, int(os.getenv("EMAIL_PORT", "993")), timeout=30)
        box.login(address, os.environ["EMAIL_PASSWORD"])
        if box.select(mailbox, readonly=True)[0] != "OK":
            raise RuntimeError(f"Cannot open mailbox {mailbox}")
        resp = box.response("UIDVALIDITY")[1]
        if not resp or not resp[0]:
            raise RuntimeError("Mailbox UIDVALIDITY is unavailable; stopping to avoid skipped or duplicate mail")
        uid_validity = resp[0].decode()
        status, found = box.uid("search", None, "SINCE", since.strftime("%d-%b-%Y"))
        if status != "OK":
            raise RuntimeError("Could not search the mailbox")
        uids = (found[0] or b"").split()
        stats["checked"] = len(uids)
        for ub in uids:
            uid = ub.decode("ascii")
            done = con.execute("SELECT status,detail FROM mail_messages WHERE mailbox=? AND uid_validity=? AND uid=?",
                               (mailbox_key, uid_validity, uid)).fetchone()
            if done and (done["status"] == "DONE" or (done["status"] == "SKIPPED" and done["detail"] == marker)):
                continue
            try:
                status, fetched = box.uid("fetch", uid, "(BODY.PEEK[])")
                raw = next((p[1] for p in fetched if isinstance(p, tuple) and isinstance(p[1], bytes)), None)
                if status != "OK" or raw is None:
                    raise RuntimeError("Could not read message")
                msg = message_from_bytes(raw)
                subject = decode_email_text(msg.get("Subject", ""))
                if flt and flt not in subject.casefold():
                    _record_message(con, mailbox_key, uid_validity, uid, "SKIPPED", subject, 0, marker)
                    stats["skipped"] += 1
                    continue
                files = [(safe_filename(decode_email_text(n)), p.get_payload(decode=True))
                         for p in msg.walk() if (n := p.get_filename())]
                supported = False
                for name, content in files:
                    if Path(name).suffix.lower() not in SUPPORTED_EXTENSIONS:
                        stats["skipped_unsupported"] += 1
                        continue
                    supported = True
                    stats["attachments"] += 1
                    for r in ingest_file(con, name, content or b"", channel="EMAIL", email_id=f"{uid_validity}:{uid}"):
                        results.append(r)
                        if r.get("route") == R_REVIEW:
                            stats["review"] += 1
                            _save_alert(con, "INTAKE_REVIEW", "Document needs intake review",
                                        f"{name}: {r.get('note', 'Check missing fields.')}", r.get("record_id") or "")
                        elif r.get("route") in (R_DUP_FILE, R_DUP_CONTENT):
                            stats["duplicates"] += 1
                        elif r.get("route") == "REJECTED":
                            _save_alert(con, "FILE_REJECTED", "Attachment could not be accepted",
                                        f"{name}: {r.get('note', 'Check file type and size.')}")
                if files and not supported:
                    _save_alert(con, "UNSUPPORTED_ATTACHMENT", "Email has no supported documents",
                                f"{subject or 'Untitled email'}: check the attachment format.")
                _record_message(con, mailbox_key, uid_validity, uid, "DONE", subject, len(files), "Processed")
                stats["emails"] += 1
            except Exception as exc:
                stats["errors"] += 1
                detail = str(exc)[:500]
                old = con.execute("SELECT status,detail FROM mail_messages WHERE mailbox=? AND uid_validity=? AND uid=?",
                                  (mailbox_key, uid_validity, uid)).fetchone()
                _record_message(con, mailbox_key, uid_validity, uid, "ERROR", "", 0, detail)
                if not old or old["status"] != "ERROR" or old["detail"] != detail:
                    _save_alert(con, "EMAIL_ERROR", "Email could not be processed", f"Message UID {uid}: {exc}")
        if stats["errors"] == 0:   # only advance the resume point after a clean scan
            con.execute("INSERT OR REPLACE INTO app_state(key,value) VALUES(?,?)", (scan_key, _now()))
        if stats["emails"] or stats["errors"]:
            _log(con, None, None, "EMAIL_CHECK", f"{stats['emails']} email(s), {stats['attachments']} attachment(s), "
                 f"{stats['skipped_unsupported']} unsupported, {stats['errors']} error(s)", "MODULE 1")
        con.commit()
        return results, stats
    finally:
        if box is not None:
            try:
                box.logout()
            except Exception:
                pass
        if own and con is not None:
            con.close()
        _POLL_LOCK.release()

def check_inbox(con):
    """Manual 'Check inbox now': one scan that waits if the background monitor is mid-scan."""
    return poll_inbox(con, blocking=True)

def _monitor_loop(stop: threading.Event):
    con = connect()   # this thread's own connection
    try:
        while not stop.is_set():
            try:
                res, stats = poll_inbox(con)
                with _MONITOR["guard"]:
                    _MONITOR.update(last=stats, last_run=_now(), last_error=None)
                    if res or stats["errors"]:
                        _MONITOR["results"] += res
                        _MONITOR["version"] += 1
            except Exception as exc:
                with _MONITOR["guard"]:
                    _MONITOR.update(last_error=str(exc)[:300], last_run=_now())
                    _MONITOR["version"] += 1
                try:
                    report_monitor_error(con, str(exc))
                except Exception:
                    pass
            stop.wait(poll_seconds())
    finally:
        con.close()

def start_monitor() -> bool:
    """Start the background scan (safe to call on every Streamlit rerun). False if already running or not configured."""
    with _MONITOR["guard"]:
        t = _MONITOR["thread"]
        if (t and t.is_alive() and not _MONITOR["stop"].is_set()) or not email_config_status()["ok"]:
            return False
        _MONITOR["stop"] = threading.Event()
        _MONITOR["thread"] = threading.Thread(target=_monitor_loop, args=(_MONITOR["stop"],), daemon=True, name="annie-inbox-monitor")
        _MONITOR["thread"].start()
        return True

def stop_monitor():
    _MONITOR["stop"].set()

def monitor_status() -> dict:
    t = _MONITOR["thread"]
    return dict(running=bool(t and t.is_alive() and not _MONITOR["stop"].is_set()), poll_seconds=poll_seconds(),
                last=_MONITOR["last"], last_run=_MONITOR["last_run"], last_error=_MONITOR["last_error"], version=_MONITOR["version"])

def pop_monitor_results() -> list[dict]:
    """Records ingested by the background scans since the last call (the UI appends them to 'Results this session')."""
    with _MONITOR["guard"]:
        out, _MONITOR["results"] = _MONITOR["results"], []
        return out

def pop_monitor_alerts() -> list[dict]:
    with _MONITOR["guard"]:
        out, _MONITOR["alerts"] = _MONITOR["alerts"], []
        return out

# =====================================================================================
# 10. MODULE 1 - INTAKE REVIEW ACTIONS
# =====================================================================================
def intake_review_rows(con) -> list[dict]:
    rows = []
    for r in con.execute("SELECT * FROM intake WHERE route=? ORDER BY received_at", (R_REVIEW,)):
        d = dict(r)
        d["extracted"], d["missing"] = json.loads(d["extracted"]), json.loads(d["missing"] or "[]")
        rows.append(d)
    return rows

def recheck_intake(con, record_id: str, edits: dict) -> dict:
    """Apply reviewer edits, re-run Module 1 validation and route again."""
    row = con.execute("SELECT * FROM intake WHERE record_id=? AND route=?", (record_id, R_REVIEW)).fetchone()
    if not row:
        raise ValueError("Record is not in Intake Review")
    raw = {**json.loads(row["extracted"]), **{k: v for k, v in edits.items() if v not in (None, "")}}
    raw.update(text_extracted=True, record_id=record_id)
    rec = clean_record(raw, row["content_fingerprint"], row["source_hash"])
    _log(con, record_id, None, "INTAKE_REVIEW_RECHECK", f"edited: {sorted(edits)} -> {rec['processing_status']}", "MODULE 1")
    ready = rec["processing_status"] == M1_READY
    con.execute("UPDATE intake SET doc_type=?, m1_status=?, route=?, extracted=?, missing=?, note=?, resolved_at=? WHERE record_id=?",
                (rec["doc_type"], rec["processing_status"], R_M2 if ready else R_REVIEW, json.dumps(rec, default=str),
                 json.dumps(rec["missing_fields"]), rec["intake_note"], _now() if ready else None, record_id))
    result = dict(record_id=record_id, m1_status=rec["processing_status"], note=rec["intake_note"],
                  missing=rec["missing_fields"], route=R_M2 if ready else R_REVIEW, source_file=row["source_file"])
    if ready:
        _log(con, record_id, None, "ROUTED_TO_MODULE_2", "READY_FOR_MODULE_2 after review", "MODULE 1")
        result["m2"] = process_document(con, to_module2_payload(rec), file_path=row["file_path"])
    con.commit()
    return result

def reextract_intake(con, record_id: str) -> dict:
    """Re-run Module 1 extraction on the stored file of an Intake Review item (use after the extraction code improves)."""
    row = con.execute("SELECT * FROM intake WHERE record_id=? AND route=?", (record_id, R_REVIEW)).fetchone()
    path = Path(row["file_path"]) if row and row["file_path"] else None
    if not path or not path.exists() or path.suffix.lower() in (".csv", ".xlsx"):
        raise ValueError("Re-extract works on stored PDF, image, DOCX or TXT files in Intake Review")
    text, method, ocr = extract_text_from_path(path)
    raw = extract_fields_from_text(path, text, method, ocr, load_reference_config(), build_reference_profiles())
    rec = clean_record({**raw, "email_id": row["email_id"], "record_id": record_id}, row["content_fingerprint"], row["source_hash"])
    con.execute("UPDATE intake SET extracted=?, extraction_method=?, ocr_used=?, classification_method=? WHERE record_id=?",
                (json.dumps(rec, default=str), method, int(ocr), rec["classification_method"], record_id))
    _log(con, record_id, None, "REEXTRACTED", f"{method}; ocr={ocr}", "MODULE 1")
    return recheck_intake(con, record_id, {})

def dismiss_intake(con, record_id: str, reason: str = ""):
    """Remove an Intake Review item entirely and forget its file hash / content fingerprint,
    so the same document can be resubmitted later without being flagged as a duplicate."""
    row = con.execute("SELECT * FROM intake WHERE record_id=? AND route=?", (record_id, R_REVIEW)).fetchone()
    if not row:
        raise ValueError("Record is not in Intake Review")
    con.execute("DELETE FROM intake WHERE record_id=?", (record_id,))
    h = row["source_hash"]
    live = con.execute("SELECT COUNT(*) FROM intake WHERE source_hash=? AND route IN (?,?)", (h, R_M2, R_REVIEW)).fetchone()[0]
    if h and live == 0:   # no other live record came from this file
        con.execute("DELETE FROM attachments WHERE source_hash=?", (h,))
        con.execute("DELETE FROM content_registry WHERE source_hash=?", (h,))
        path = Path(row["file_path"]) if row["file_path"] else None
        if path and path.exists() and FILE_DIR in path.parents:
            path.unlink()
    _log(con, record_id, None, "INTAKE_DISMISSED", f"{row['source_file']} removed; can be resubmitted. {reason}".strip(), "MODULE 1")
    con.commit()

# =====================================================================================
# 11. MODULE 2 - SIMILARITY HELPERS AND LINE SCORING
# =====================================================================================
def _norm(s) -> str:
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9 ]", "", str(s or "").lower())).strip()

def _ref_norm(s) -> str:
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())

def text_sim(a, b):
    a, b = _norm(a), _norm(b)
    if not a or not b:
        return None
    ta, tb = set(a.split()), set(b.split())
    return max(SequenceMatcher(None, a, b).ratio(), len(ta & tb) / len(ta | tb))

def ref_sim(a, b):
    a, b = _ref_norm(a), _ref_norm(b)
    return SequenceMatcher(None, a, b).ratio() if a and b else None

def num_sim(a, b):
    try:
        a, b = float(a), float(b)
    except (TypeError, ValueError):
        return None
    if a == b:
        return 1.0
    return max(0.0, 1 - abs(a - b) / max(abs(a), abs(b)) * 10)   # a 10% gap scores 0

def parse_date(s):
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d %b %Y"):
        try:
            return datetime.strptime(str(s), fmt).date()
        except ValueError:
            pass
    return None

def doc_total(d):
    if d.get("total") not in (None, ""):
        return d["total"]
    try:
        items = d.get("items") or []
        return round(sum(float(i.get("total", 0)) for i in items), 2) if items else None
    except (TypeError, ValueError):
        return None

def weighted(scores: dict, weights: dict) -> float:
    use = {k: s for k, s in scores.items() if s is not None}   # missing fields do not penalise
    w = sum(weights[k] for k in use)
    return round(sum(weights[k] * s for k, s in use.items()) / w, 4) if w else 0.0

LINE_W = dict(desc=0.6, price=0.4)    # invoice line vs PO line
VENDOR_MIN = 0.7                      # vendor names must be at least this similar (tightened from 0.6)
CITATION_BONUS = 0.25                 # added when the invoice cites the PO the line belongs to
ENFORCE_UNIQUE_INVOICE_NUMBER = True  # exact vendor + invoice number may only be allocated once (not fuzzy)

def line_score(inv_line: dict, po_line: dict, cited: bool, sole_open_line: bool) -> float:
    """Confidence that an invoice line belongs to a PO line. One threshold (LINK_THRESHOLD) decides."""
    if cited and sole_open_line:
        return 1.0   # the cited PO has exactly one line left to bill
    desc = text_sim(inv_line.get("description"), po_line.get("description"))
    price = num_sim(inv_line.get("unit_price"), po_line.get("unit_price")) if po_line["basis"] == "QTY" else None
    s = weighted(dict(desc=desc, price=price), LINE_W)
    return round(min(1.0, s + CITATION_BONUS) if cited else s, 4)

# =====================================================================================
# 12. MODULE 2 - VALIDATION HELPERS
# =====================================================================================
def normalise_type(t) -> str:
    t = str(t or "").upper().replace(" ", "_").replace("/", "_")
    if t in ("PO", "PR", "PR_PO", "PURCHASE_REQUISITION", "PURCHASE_ORDER", "APPROVAL_FORM"):
        return PO
    if t in ("INVOICE", "INV", "SUPPLIER_INVOICE"):
        return INV
    if t in ("GR", "SES", "GR_SES"):
        return GR
    return t

def _blank(v) -> bool:
    return v is None or str(v).strip() == ""

def _cites_amount(d) -> bool:
    """Does the document cite any amount (a total, or a unit price / line total on any item)?"""
    try:
        if float(d.get("total") or 0) > 0:
            return True
    except (TypeError, ValueError):
        pass
    return any(not _blank(i.get("unit_price")) or not _blank(i.get("total")) for i in d.get("items") or [])

def missing_fields(d) -> list[str]:
    """Mandatory-field check used by the Module 1 gate."""
    out = []
    for f in MANDATORY[d["doc_type"]]:
        v = d.get(f)
        if f == "signature_present":
            if v not in (True, "Yes", "yes", "Y", "true", "True"):
                out.append(f)
        elif f == "lines":
            cat = str(d.get("category", "")).lower()
            items = d.get("items") or []
            if d["doc_type"] == PO and cat.startswith("good") and not _cites_amount(d):
                has_goods = bool(items) and all(not _blank(i.get("description")) for i in items)   # call-off PO/PR: no amounts needed
            else:
                has_goods = bool(items) and all(not _blank(i.get("description")) and not _blank(i.get("qty"))
                                                and not _blank(i.get("unit_price")) and not _blank(i.get("total")) for i in items)
            has_service = not _blank(d.get("service_description")) and str(d.get("service_description")).upper() != "N/A"
            ok = has_goods if cat.startswith("good") else has_service if cat.startswith("serv") else (has_goods or has_service)
            if not ok:
                out.append(f)
        elif f == "total":
            if doc_total(d) in (None, ""):
                out.append(f)
        elif _blank(v):
            out.append(f)
    return sorted(set(out))

def get_doc(con, rid):
    if not rid:
        return None
    r = con.execute("SELECT * FROM documents WHERE record_id=?", (rid,)).fetchone()
    if not r:
        return None
    d = dict(r)
    d["payload"], d["missing"] = json.loads(d["payload"]), json.loads(d["missing"] or "[]")
    return d

def _new_case_id(con) -> str:
    """Next serial for this year. Uses the highest number in use (not a row count) so deleting a case never causes a clash."""
    year = date.today().year
    used = [int(m.group(1)) for r in con.execute("SELECT case_id FROM cases") if (m := re.match(rf"^CASE-{year}-(\d+)$", r[0]))]
    return f"CASE-{year}-{(max(used) if used else 0) + 1:04d}"

# =====================================================================================
# 13. MODULE 2 - PO REGISTRATION, INVOICE ALLOCATION AND CASE STATUS
# =====================================================================================
def _po_lines(po: dict) -> list[dict]:
    """PO/PR -> billable lines. Goods: by quantity. Services: by amount if one is cited, else SINGLE (one invoice only)."""
    items = po.get("items") or []
    if str(po.get("category", "")).lower().startswith("good") and items:
        return [dict(description=i.get("description"), basis="QTY", ordered_qty=float(i["qty"]),
                     unit_price=float(i["unit_price"]), ordered_amount=float(i["total"])) for i in items]
    total = float(doc_total(po) or 0)
    if total > 0:   # services that cite an amount: tracked by amount (many invoices allowed)
        return [dict(description=po.get("service_description") or "", basis="AMOUNT", ordered_qty=1.0,
                     unit_price=total, ordered_amount=total)]
    # no amount cited: the PO/PR expects exactly ONE invoice (one-to-one)
    return [dict(description=po.get("service_description") or "", basis="SINGLE", ordered_qty=1.0,
                 unit_price=0.0, ordered_amount=0.0)]

def _invoice_lines(p: dict) -> list[dict]:
    items = p.get("items") or []
    if items:
        return [dict(description=i.get("description"), qty=float(i.get("qty") or 0), unit_price=float(i.get("unit_price") or 0),
                     amount=float(i.get("total") or 0), po=_ref_norm(i.get("po_reference"))) for i in items]
    total = float(doc_total(p) or 0)
    return [dict(description=p.get("service_description"), qty=1.0, unit_price=total, amount=total, po="")]

def _cited(p: dict) -> dict:
    refs = list(p.get("po_references") or []) + ([p["po_reference"]] if p.get("po_reference") else [])
    refs += [i["po_reference"] for i in (p.get("items") or []) if i.get("po_reference")]
    return {_ref_norm(r): r for r in refs if r}

def _need(ln: dict, l: dict) -> float:
    """How much of a PO line an invoice line consumes: qty (goods), 1 (single-invoice service), amount (services by amount)."""
    return ln["qty"] if l["basis"] == "QTY" else (1.0 if l["basis"] == "SINGLE" else ln["amount"])

def _alloc_qty(ln: dict, l: dict) -> float:
    return ln["qty"] if l["basis"] == "QTY" else (1.0 if l["basis"] == "SINGLE" else 0.0)

def _open(l: dict) -> float:
    return (l["ordered_qty"] - l["billed_qty"]) if l["basis"] in ("QTY", "SINGLE") else (l["ordered_amount"] - l["billed_amount"])

def _new_calloff_id(con) -> str:
    year = date.today().year
    used = [int(m.group(1)) for r in con.execute("SELECT case_id FROM calloffs") if (m := re.match(rf"^CALL-{year}-(\d+)$", r[0]))]
    return f"CALL-{year}-{(max(used) if used else 0) + 1:04d}"

def _calloff_refs(con) -> dict:
    """PO/PR numbers that are Call-Off cases (excluded from automatic matching)."""
    out = {}
    for r in con.execute("SELECT d.payload FROM calloffs c JOIN documents d ON d.record_id=c.po_record_id"):
        ref = json.loads(r[0]).get("po_reference")
        if ref:
            out[_ref_norm(ref)] = ref
    return out

def _open_calloff(con, doc: dict, rid: str, file_path) -> dict:
    case_id, now = _new_calloff_id(con), _now()
    con.execute("INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?)",
                (rid, PO, json.dumps(doc), file_path, now, "CALLOFF", case_id, None, None, "[]"))
    con.execute("INSERT INTO calloffs(case_id, po_record_id, status, reason, created_at, updated_at) VALUES(?,?,?,?,?,?)",
                (case_id, rid, S_CALLOFF, "", now, now))
    _log(con, rid, case_id, "CALLOFF_OPENED", f"{doc.get('po_reference')} cites no amount: manual matching only")
    retry_unallocated(con)   # waiting invoices that cite this PO/PR now get the right explanation
    con.commit()
    return dict(record_id=rid, doc_type=PO, outcome="CALLOFF_OPENED", case_id=case_id, status=S_CALLOFF, po_reference=doc.get("po_reference"))

def register_po(con, doc: dict, rid: str, file_path) -> dict:
    ref = doc.get("po_reference")
    for r in con.execute("SELECT record_id, payload FROM documents WHERE doc_type=?", (PO,)):
        if _ref_norm(json.loads(r["payload"]).get("po_reference")) == _ref_norm(ref):
            _log(con, rid, None, "REJECTED_PO_REFERENCE_EXISTS", f"{ref} already registered as {r['record_id']}")
            con.commit()
            return dict(record_id=rid, doc_type=PO, outcome="REJECTED_PO_REFERENCE_EXISTS", detail=f"{ref} already registered")
    if not _cites_amount(doc):   # no amount cited: a Call-Off case, matched manually to several invoices
        return _open_calloff(con, doc, rid, file_path)
    con.execute("INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?)",
                (rid, PO, json.dumps(doc), file_path, _now(), "REGISTERED", None, None, None, "[]"))
    lines = _po_lines(doc)
    for n, ln in enumerate(lines, 1):
        con.execute("INSERT INTO po_lines(po_record_id, po_reference, line_no, description, basis, ordered_qty, unit_price, ordered_amount)"
                    " VALUES(?,?,?,?,?,?,?,?)", (rid, ref, n, ln["description"], ln["basis"], ln["ordered_qty"], ln["unit_price"], ln["ordered_amount"]))
    _log(con, rid, None, "PO_REGISTERED", f"{ref}: {len(lines)} line(s), total {doc_total(doc)}")
    newly = retry_unallocated(con)
    con.commit()
    return dict(record_id=rid, doc_type=PO, outcome="REGISTERED", po_reference=ref, lines=len(lines),
                allocated_waiting_invoices=[(r["case_id"], r.get("invoice_reference")) for r in newly])

def allocate_invoice(con, case_id: str, force: bool = False) -> dict:
    """All-or-nothing: every invoice line must find a PO line with enough open balance, else the invoice waits (Stage 1)."""
    case = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not case or case["status"] != S_OPEN_PO:
        return dict(case_id=case_id, allocated=False, reason="case is not waiting for allocation")
    if case["hold"] and not force:
        return dict(case_id=case_id, allocated=False, reason="held after a manual unmatch", status=S_OPEN_PO)
    if force and case["hold"]:
        con.execute("UPDATE cases SET hold=0 WHERE case_id=?", (case_id,))
    inv = get_doc(con, case["inv_record_id"])
    p, lines, cited = inv["payload"], _invoice_lines(inv["payload"]), _cited(inv["payload"])
    po_lines = []
    for r in con.execute("SELECT l.*, d.payload AS po_payload FROM po_lines l JOIN documents d ON d.record_id=l.po_record_id ORDER BY l.line_id"):
        d = dict(r)
        d["vendor"] = json.loads(d.pop("po_payload")).get("vendor_name")
        po_lines.append(d)
    reasons, plan = [], []
    known = {_ref_norm(l["po_reference"]) for l in po_lines}
    calloff = _calloff_refs(con)
    absent = [orig for n, orig in cited.items() if n not in known and n not in calloff]
    via_calloff = [orig for n, orig in cited.items() if n in calloff]
    if absent or via_calloff:
        if absent:
            reasons.append("PO/PR not yet received: " + ", ".join(absent))
        if via_calloff:
            reasons.append(f"{', '.join(via_calloff)} is a Call-Off PO/PR (no amount cited): match it manually from the Call-Offs tab")
    else:
        local = {l["line_id"]: _open(l) for l in po_lines}
        for ln in lines:
            allowed = {ln["po"]} if ln["po"] else (set(cited) if cited else None)
            scored = []
            for l in po_lines:
                if (text_sim(p.get("vendor_name"), l["vendor"]) or 0) < VENDOR_MIN or local[l["line_id"]] <= 1e-9:
                    continue
                if allowed is not None and _ref_norm(l["po_reference"]) not in allowed:
                    continue
                is_cited = _ref_norm(l["po_reference"]) in cited or ln["po"] == _ref_norm(l["po_reference"])
                sole = sum(1 for x in po_lines if x["po_record_id"] == l["po_record_id"] and local[x["line_id"]] > 1e-9) == 1
                scored.append((line_score(ln, l, is_cited, sole and is_cited), l))
            scored.sort(key=lambda t: (-t[0], t[1]["line_id"]))
            good = [t for t in scored if t[0] >= LINK_THRESHOLD]
            label = (ln["description"] or "")[:40]
            if not good:
                if scored:
                    why = f" (best {scored[0][0]:.2f})"
                else:
                    spent = sorted({l["po_reference"] for l in po_lines if (allowed is None or _ref_norm(l["po_reference"]) in allowed)
                                    and (text_sim(p.get("vendor_name"), l["vendor"]) or 0) >= VENDOR_MIN})
                    why = f": {', '.join(spent)} has nothing left to bill" if spent else " (no PO/PR for this vendor)"
                reasons.append(f"'{label}' matches no open PO line" + why)
                continue
            pick = None
            for s, l in good:
                need = _need(ln, l)
                if need <= local[l["line_id"]] + (1e-9 if l["basis"] in ("QTY", "SINGLE") else 0.005):
                    pick = (s, l, need)
                    break
            if not pick:
                s, l = good[0]
                need = _need(ln, l)
                reasons.append(f"'{label}' bills {need:g} but {l['po_reference']} has only {local[l['line_id']]:g} left to bill")
                continue
            s, l, need = pick
            local[l["line_id"]] -= need
            plan.append(dict(line=l, qty=_alloc_qty(ln, l), amount=ln["amount"], score=s))
    now = _now()
    if reasons or not plan:
        reason = "; ".join(reasons) or "no lines to allocate"
        if reason != (case["reason"] or ""):
            _log(con, inv["record_id"], case_id, "ALLOCATION_PENDING", reason)
        con.execute("UPDATE cases SET reason=?, updated_at=? WHERE case_id=?", (reason, now, case_id))
        con.commit()
        return dict(case_id=case_id, allocated=False, reason=reason, status=S_OPEN_PO)
    for a in plan:
        con.execute("INSERT INTO allocations(doc_record_id, case_id, kind, po_line_id, qty, amount, score, created_at) VALUES(?,?,?,?,?,?,?,?)",
                    (inv["record_id"], case_id, INV, a["line"]["line_id"], a["qty"], a["amount"], a["score"], now))
        con.execute("UPDATE po_lines SET billed_qty=billed_qty+?, billed_amount=billed_amount+? WHERE line_id=?",
                    (a["qty"], a["amount"], a["line"]["line_id"]))
    score = round(sum(a["score"] for a in plan) / len(plan), 4)
    con.execute("UPDATE cases SET reason='', link_score=?, updated_at=? WHERE case_id=?", (score, now, case_id))
    con.execute("UPDATE documents SET outcome='ALLOCATED' WHERE record_id=?", (inv["record_id"],))
    status = refresh_case(con, case_id)
    pos = sorted({a["line"]["po_reference"] for a in plan})
    _log(con, inv["record_id"], case_id, "ALLOCATED", f"to {', '.join(pos)}; score {score}; status={status}")
    con.commit()
    return dict(case_id=case_id, allocated=True, status=status, score=score, po_references=pos,
                invoice_reference=p.get("invoice_reference"), reason="")

def retry_unallocated(con) -> list[dict]:
    """Re-try every waiting invoice (oldest first), e.g. after a new PO/PR arrives."""
    out = []
    for c in con.execute("SELECT case_id FROM cases WHERE status=? AND COALESCE(hold,0)=0 ORDER BY case_id", (S_OPEN_PO,)).fetchall():
        r = allocate_invoice(con, c["case_id"])
        if r.get("allocated"):
            out.append(r)
    return out

def _case_receipts(con, case_id: str) -> dict:
    """GR/SES receipts raised against this invoice's match: line_id -> [qty, amount]."""
    out = {}
    for r in con.execute("SELECT payload FROM documents WHERE doc_type=? AND case_id=? AND outcome='GR_SUBMITTED'", (GR, case_id)):
        for x in json.loads(r["payload"]).get("lines", []):
            q, a_ = out.setdefault(int(x["line_id"]), [0.0, 0.0])
            out[int(x["line_id"])] = [q + float(x.get("qty") or 0), a_ + float(x.get("amount") or 0)]
    return out

def refresh_case(con, case_id: str) -> str:
    """Pending GR/SES until this invoice's own receipts cover what it billed on every allocated PO line."""
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if c["status"] in (S_ARCHIVED, S_CALLED):
        return c["status"]
    allocs = con.execute("SELECT * FROM allocations WHERE case_id=? ORDER BY alloc_id", (case_id,)).fetchall()
    if not allocs:
        status = S_OPEN_PO
    else:
        need = {}
        for a in allocs:
            n = need.setdefault(a["po_line_id"], [0.0, 0.0])
            n[0] += a["qty"]
            n[1] += a["amount"]
        got = _case_receipts(con, case_id)
        covered = True
        for lid, (nq, na) in need.items():
            l = con.execute("SELECT basis FROM po_lines WHERE line_id=?", (lid,)).fetchone()
            gq, ga = got.get(lid, [0.0, 0.0])
            covered &= (gq + 1e-9 >= nq) if l["basis"] in ("QTY", "SINGLE") else (ga + 0.005 >= na)
        status = S_DONE if covered else S_GR
    con.execute("UPDATE cases SET status=?, updated_at=? WHERE case_id=?", (status, _now(), case_id))
    return status

def refresh_all(con):
    for c in con.execute("SELECT case_id FROM cases WHERE status IN (?,?)", (S_GR, S_DONE)).fetchall():
        refresh_case(con, c["case_id"])

def process_document(con, doc: dict, file_path=None) -> dict:
    """Module 2 entry point. PO/PR -> registered with line balances; invoice -> new case + allocation attempt.
    Duplicate and completeness checks are done in Module 1."""
    doc = dict(doc)
    doc["doc_type"] = normalise_type(doc.get("doc_type"))
    rid = str(doc.get("record_id") or "").strip()
    if doc["doc_type"] not in (PO, INV):
        raise ValueError("doc_type must be PO (PR/PO) or INVOICE. GR/SES is submitted from the case page.")
    if not rid:
        raise ValueError("Document is missing 'record_id'")
    if con.execute("SELECT 1 FROM documents WHERE record_id=?", (rid,)).fetchone():
        return dict(record_id=rid, outcome="ALREADY_PROCESSED", detail="record_id already in Module 2")
    if doc["doc_type"] == PO:
        return register_po(con, doc, rid, file_path)

    if ENFORCE_UNIQUE_INVOICE_NUMBER:   # exact, not fuzzy: the same invoice must never consume PO balance twice
        for r in con.execute("SELECT record_id, payload FROM documents WHERE doc_type=?", (INV,)):
            o = json.loads(r["payload"])
            if _ref_norm(o.get("invoice_reference")) == _ref_norm(doc.get("invoice_reference")) \
                    and (text_sim(o.get("vendor_name"), doc.get("vendor_name")) or 0) >= 0.9:
                _log(con, rid, None, "REJECTED_INVOICE_NUMBER_EXISTS", f"{doc.get('invoice_reference')} already held as {r['record_id']}")
                con.commit()
                return dict(record_id=rid, doc_type=INV, outcome="REJECTED_INVOICE_NUMBER_EXISTS",
                            detail=f"{doc.get('invoice_reference')} from this vendor is already in Annie ({r['record_id']})")
    case_id, now = _new_case_id(con), _now()
    con.execute("INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?)",
                (rid, INV, json.dumps(doc), file_path, now, "UNALLOCATED", case_id, None, None, "[]"))
    con.execute("INSERT INTO cases(case_id, inv_record_id, status, reason, link_score, created_at, updated_at) VALUES(?,?,?,?,?,?,?)",
                (case_id, rid, S_OPEN_PO, "", None, now, now))
    _log(con, rid, case_id, "INVOICE_RECEIVED", f"{doc.get('invoice_reference')} total {doc_total(doc)}")
    res = allocate_invoice(con, case_id)
    con.commit()
    return dict(record_id=rid, doc_type=INV, outcome="ALLOCATED" if res["allocated"] else "UNALLOCATED", case_id=case_id,
                status=res.get("status"), score=res.get("score"), po_references=res.get("po_references"), reason=res.get("reason"))

def load_samples(con, folder: Path | str) -> list[dict]:
    """Test loader: pre-extracted sample JSONs stand in for Module 1 text extraction (duplicate checks and gate still apply)."""
    out = []
    for jf in sorted(Path(folder).glob("*.json")):
        payload = json.loads(jf.read_text())
        pdf, stored, h, fp = jf.with_suffix(".pdf"), None, None, None
        if pdf.exists():
            data = pdf.read_bytes()
            h = calculate_hash(data)
            seen = con.execute("SELECT * FROM attachments WHERE source_hash=?", (h,)).fetchone()
            if seen:
                out.append(_dup_result(con, pdf.name, "SAMPLE_JSON", None, h, "DUPLICATE_FILE", seen["original_filename"], seen["stored_path"]))
                continue
            FILE_DIR.mkdir(parents=True, exist_ok=True)
            stored = FILE_DIR / f"{h[:10]}_{safe_filename(pdf.name)}"
            stored.write_bytes(data)
            con.execute("INSERT INTO attachments VALUES(?,?,?,?,?)", (h, str(stored), None, pdf.name, _now()))
        rec = clean_record({**payload, "text_extracted": True, "source_file": pdf.name if stored else jf.name,
                            "extraction_method": "SAMPLE_JSON", "classification_method": "SAMPLE_JSON"}, fp, h)
        out.append(route_record(con, rec, "SAMPLE_JSON", stored))
    return out

# =====================================================================================
# 14. MODULE 2 - CASE DETAIL AND GR/SES BATCHES
# =====================================================================================
def case_lines(con, case_id: str) -> list[dict]:
    """Allocated lines of an invoice case (one row per PO line) with what this invoice billed and what has been received for it."""
    rows = con.execute("SELECT a.alloc_id, a.qty AS a_qty, a.amount AS a_amt, a.score, a.manual, a.reason AS a_reason, l.* FROM allocations a "
                       "JOIN po_lines l ON l.line_id=a.po_line_id WHERE a.case_id=? ORDER BY a.alloc_id", (case_id,)).fetchall()
    agg = {}
    for r in rows:
        d = dict(r)
        if d["line_id"] in agg:
            agg[d["line_id"]]["a_qty"] += d["a_qty"]
            agg[d["line_id"]]["a_amt"] += d["a_amt"]
        else:
            agg[d["line_id"]] = d
    got, out = _case_receipts(con, case_id), []
    for d in agg.values():
        qty = d["basis"] in ("QTY", "SINGLE")
        ordered, po_received = (d["ordered_qty"], d["received_qty"]) if qty else (d["ordered_amount"], d["received_amount"])
        billed = d["a_qty"] if qty else d["a_amt"]
        received = got.get(d["line_id"], [0.0, 0.0])[0 if qty else 1]
        d.update(billed_this=billed, billed_cum=billed, ordered=ordered, received=received, po_received=po_received,
                 needed=round(max(0.0, billed - received), 4), room=round(max(0.0, ordered - po_received), 4))
        out.append(d)
    return out

def case_po_docs(con, case_id: str) -> list[dict]:
    ids = [r[0] for r in con.execute("SELECT DISTINCT l.po_record_id FROM allocations a JOIN po_lines l ON l.line_id=a.po_line_id "
                                     "WHERE a.case_id=? ORDER BY l.po_record_id", (case_id,))]
    return [get_doc(con, i) for i in ids]

def case_gr_docs(con, case_id: str) -> list[dict]:
    """GR/SES batches raised for THIS invoice-PO match only (not receipts of other invoices on the same PO)."""
    return [get_doc(con, r[0]) for r in con.execute(
        "SELECT record_id FROM documents WHERE doc_type=? AND case_id=? AND outcome='GR_SUBMITTED' ORDER BY received_at, record_id", (GR, case_id))]

def next_ses_number(con, category) -> str:
    """Serial number for the next GR (goods) or SES (services), e.g. GR-2026-0001. Never typed by the user."""
    prefix = "GR" if str(category or "").lower().startswith("good") else "SES"
    year = date.today().year
    pat = re.compile(rf"^{prefix}-{year}-(\d+)$")
    used = [int(m.group(1)) for r in con.execute("SELECT payload FROM documents WHERE doc_type=?", (GR,))
            if (m := pat.match(str(json.loads(r["payload"]).get("ses_gr_number", ""))))]
    return f"{prefix}-{year}-{(max(used) if used else 0) + 1:04d}"

def allowed_preparers(con, case_id: str) -> list[str]:
    return sorted({d["payload"].get("prepared_by") for d in case_po_docs(con, case_id) if d["payload"].get("prepared_by")})

def submit_gr(con, case_id: str, gr: dict) -> dict:
    """Record one GR/SES batch. gr = {service_date_from, service_date_to, prepared_by, lines:[{line_id, receive}]}.
    'receive' is a quantity for goods lines and an amount for service lines."""
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] != S_GR:
        raise ValueError("Case is not pending GR/SES")
    lines = {l["line_id"]: l for l in case_lines(con, case_id)}
    wanted = [(int(x["line_id"]), float(x.get("receive") or 0)) for x in gr.get("lines", []) if float(x.get("receive") or 0) > 0]
    if not wanted:
        raise ValueError("Enter a received quantity or amount for at least one line")
    if _norm(gr.get("prepared_by")) not in {_norm(n) for n in allowed_preparers(con, case_id)}:
        raise ValueError("Prepared by must match the PO/PR: " + " / ".join(allowed_preparers(con, case_id)))
    d1, d2 = parse_date(gr.get("service_date_from")), parse_date(gr.get("service_date_to"))
    if not d1 or not d2:
        raise ValueError("Performance / service dates are required")
    if d2 < d1:
        raise ValueError("Performance end date is before start date")
    out_lines, total = [], 0.0
    for lid, recv in wanted:
        if lid not in lines:
            raise ValueError("Line does not belong to this case")
        l = lines[lid]
        cap = min(l["needed"], l["room"])
        if recv > cap + (1e-9 if l["basis"] in ("QTY", "SINGLE") else 0.005):
            raise ValueError(f"{l['po_reference']} '{(l['description'] or '')[:30]}': receiving {recv:g} exceeds the {cap:g} this invoice still needs "
                             f"(and the PO/PR has left to receive)")
        amount = (round(recv * l["unit_price"], 2) if l["basis"] == "QTY"
                  else round(l["a_amt"] * min(recv, 1.0), 2) if l["basis"] == "SINGLE" else round(recv, 2))
        total += amount
        out_lines.append(dict(line_id=lid, po_reference=l["po_reference"], description=l["description"], basis=l["basis"],
                              qty=recv if l["basis"] in ("QTY", "SINGLE") else None, amount=amount))
    category = "Services" if all(x["basis"] in ("AMOUNT", "SINGLE") for x in out_lines) else "Goods"
    number = next_ses_number(con, category)
    po_vendor = case_po_docs(con, case_id)[0]["payload"].get("vendor_name")
    payload = dict(doc_type=GR, record_id=number, ses_gr_number=number, case_id=case_id, vendor_name=po_vendor, category=category,
                   po_references=sorted({x["po_reference"] for x in out_lines}), service_date_from=gr["service_date_from"],
                   service_date_to=gr["service_date_to"], prepared_by=gr["prepared_by"], lines=out_lines, total=round(total, 2))
    con.execute("INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?)",
                (number, GR, json.dumps(payload), None, _now(), "GR_SUBMITTED", case_id, None, None, "[]"))
    for x in out_lines:
        if x["basis"] in ("QTY", "SINGLE"):
            con.execute("UPDATE po_lines SET received_qty=received_qty+?, received_amount=received_amount+? WHERE line_id=?",
                        (x["qty"], x["amount"], x["line_id"]))
        else:
            con.execute("UPDATE po_lines SET received_amount=received_amount+? WHERE line_id=?", (x["amount"], x["line_id"]))
    refresh_all(con)
    status = con.execute("SELECT status FROM cases WHERE case_id=?", (case_id,)).fetchone()[0]
    _log(con, number, case_id, "GR_SUBMITTED", f"{number}: {len(out_lines)} line(s), total {payload['total']}; case status={status}")
    con.commit()
    return dict(record_id=number, case_id=case_id, status=status, ses_gr_number=number)

# =====================================================================================
# 14C. MANUAL MATCHING, UNMATCH AND INVOICE DISMISSAL
# =====================================================================================
def _reverse_allocations(con, case_id: str) -> int:
    """Give back everything an invoice case billed against PO lines (works for qty, amount and single-invoice lines)."""
    rows = con.execute("SELECT * FROM allocations WHERE case_id=?", (case_id,)).fetchall()
    for a in rows:
        con.execute("UPDATE po_lines SET billed_qty=MAX(0, billed_qty-?), billed_amount=MAX(0, billed_amount-?) WHERE line_id=?",
                    (a["qty"], a["amount"], a["po_line_id"]))
    con.execute("DELETE FROM allocations WHERE case_id=?", (case_id,))
    return len(rows)

def _void_case_receipts(con, case_id: str) -> int:
    """A match was undone, so the GR/SES batches raised for it are voided and their received quantities given back."""
    n = 0
    for r in con.execute("SELECT record_id, payload FROM documents WHERE doc_type=? AND case_id=? AND outcome='GR_SUBMITTED'", (GR, case_id)).fetchall():
        for x in json.loads(r["payload"]).get("lines", []):
            con.execute("UPDATE po_lines SET received_qty=MAX(0, received_qty-?), received_amount=MAX(0, received_amount-?) WHERE line_id=?",
                        (float(x.get("qty") or 0), float(x.get("amount") or 0), x["line_id"]))
        con.execute("UPDATE documents SET outcome='GR_VOIDED' WHERE record_id=?", (r["record_id"],))
        _log(con, r["record_id"], case_id, "GR_VOIDED", "match reversed")
        n += 1
    return n

def delete_invoice_case(con, case_id: str, reason: str = "") -> dict:
    """Dismiss (delete) an invoice in Stage 1 or Stage 2. Its billing is reversed on the PO/PR lines (a single-invoice
    PO/PR becomes open again), the file hash is forgotten so the invoice can be resubmitted later, and the audit log keeps a line.
    GR/SES batches raised for this invoice-PO match are voided with it (their received quantities are given back)."""
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] not in (S_OPEN_PO, S_GR):
        raise ValueError("Only invoices in Stage 1 or Stage 2 can be dismissed")
    inv = get_doc(con, c["inv_record_id"])
    _void_case_receipts(con, case_id)
    reversed_n = _reverse_allocations(con, case_id)
    con.execute("DELETE FROM cases WHERE case_id=?", (case_id,))
    con.execute("DELETE FROM documents WHERE record_id=?", (inv["record_id"],))
    row = con.execute("SELECT * FROM intake WHERE record_id=?", (inv["record_id"],)).fetchone()
    if row:
        con.execute("DELETE FROM intake WHERE record_id=?", (inv["record_id"],))
        h = row["source_hash"]
        live = con.execute("SELECT COUNT(*) FROM intake WHERE source_hash=? AND route IN (?,?)", (h, R_M2, R_REVIEW)).fetchone()[0]
        if h and live == 0:
            con.execute("DELETE FROM attachments WHERE source_hash=?", (h,))
            con.execute("DELETE FROM content_registry WHERE source_hash=?", (h,))
            path = Path(row["file_path"]) if row["file_path"] else None
            if path and path.exists() and FILE_DIR in path.parents:
                path.unlink()
    _log(con, inv["record_id"], case_id, "INVOICE_DISMISSED",
         f"{inv['payload'].get('invoice_reference')} deleted; {reversed_n} allocation(s) reversed. {reason}".strip())
    refresh_all(con)
    retry_unallocated(con)
    con.commit()
    return dict(case_id=case_id, reversed_allocations=reversed_n)

def unmatch_case(con, case_id: str, reason: str) -> dict:
    """Stage 2 -> Stage 1: undo the match but keep the invoice. The case is held, so it is not re-matched automatically."""
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] != S_GR:
        raise ValueError("Only invoices in Stage 2 (Pending GR/SES) can be unmatched")
    if len((reason or "").strip()) < 5:
        raise ValueError("A typed reason (at least 5 characters) is required to unmatch")
    _void_case_receipts(con, case_id)
    n = _reverse_allocations(con, case_id)
    con.execute("UPDATE cases SET status=?, reason=?, link_score=NULL, hold=1, updated_at=? WHERE case_id=?",
                (S_OPEN_PO, f"Unmatched by user: {reason.strip()}", _now(), case_id))
    con.execute("UPDATE documents SET outcome='UNALLOCATED' WHERE record_id=?", (c["inv_record_id"],))
    _log(con, c["inv_record_id"], case_id, "UNMATCHED", f"{n} allocation(s) reversed. {reason.strip()}")
    refresh_all(con)
    retry_unallocated(con)
    con.commit()
    return dict(case_id=case_id, reversed_allocations=n, status=S_OPEN_PO)

def invoice_line_list(con, case_id: str) -> list[dict]:
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    p = get_doc(con, c["inv_record_id"])["payload"]
    return [dict(index=i, **ln) for i, ln in enumerate(_invoice_lines(p), 1)]

def _po_line_universe(con) -> dict:
    out = {}
    for r in con.execute("SELECT l.*, d.payload AS po_payload FROM po_lines l JOIN documents d ON d.record_id=l.po_record_id ORDER BY l.line_id"):
        d = dict(r)
        d["vendor"] = json.loads(d.pop("po_payload")).get("vendor_name")
        out[d["line_id"]] = d
    return out

def manual_candidates(con, case_id: str) -> list[dict]:
    """PO/PR lines that still have something left to bill, most similar vendor first (for the manual-match picker)."""
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    vendor = get_doc(con, c["inv_record_id"])["payload"].get("vendor_name")
    unit = {"QTY": "qty", "AMOUNT": "amount", "SINGLE": "single invoice"}
    out = []
    for l in _po_line_universe(con).values():
        left = _open(l)
        if left <= (1e-9 if l["basis"] in ("QTY", "SINGLE") else 0.005):
            continue
        vs = text_sim(vendor, l["vendor"]) or 0.0
        out.append(dict(line_id=l["line_id"], po_reference=l["po_reference"], vendor=l["vendor"], description=l["description"],
                        basis=l["basis"], open=left, unit_price=l["unit_price"], vendor_sim=round(vs, 3), line_no=l["line_no"],
                        label=f"#{l['line_id']} {l['po_reference']} | {(l['description'] or '')[:45]} | left {left:g} {unit[l['basis']]} | {l['vendor']}"))
    out.sort(key=lambda x: (-x["vendor_sim"], x["po_reference"], x["line_no"]))
    return out

def manual_match(con, case_id: str, rows: list[dict], reason: str = "", dry_run: bool = False) -> dict:
    """Match an invoice by hand. rows = [{inv_line, po_line_id, value}] where value is a quantity (goods line), an amount
    (service by amount) or any positive number (single-invoice service: it takes the whole invoice line).
    One invoice line can be split over several PO lines; every invoice line must end up 100% allocated.
    Capacity limits and the one-invoice rule still apply. dry_run only reports scores, warnings and errors."""
    case = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not case or case["status"] != S_OPEN_PO:
        raise ValueError("Only invoices waiting in Stage 1 can be matched manually")
    p = get_doc(con, case["inv_record_id"])["payload"]
    lines, cited = _invoice_lines(p), _cited(p)
    pol = _po_line_universe(con)
    local = {i: _open(l) for i, l in pol.items()}
    open_per_po = {}
    for l in pol.values():
        if _open(l) > 1e-9:
            open_per_po[l["po_record_id"]] = open_per_po.get(l["po_record_id"], 0) + 1
    errors, out, cover, single_used = [], [], {i: 0.0 for i in range(1, len(lines) + 1)}, set()
    for r in rows:
        try:
            i, pid, v = int(r["inv_line"]), int(r["po_line_id"]), float(r["value"] or 0)
        except (TypeError, ValueError, KeyError):
            errors.append("A row is incomplete (needs invoice line, PO/PR line and a number)")
            continue
        if i < 1 or i > len(lines):
            errors.append(f"Invoice line {i} does not exist")
            continue
        if pid not in pol:
            errors.append(f"Row for invoice line {i}: PO/PR line not found")
            continue
        if v <= 0:
            errors.append(f"Invoice line {i}: enter a quantity or amount above zero")
            continue
        ln, l = lines[i - 1], pol[pid]
        label = f"line {i} -> {l['po_reference']} '{(l['description'] or '')[:30]}'"
        if l["basis"] == "QTY":
            need, qty, frac = v, v, (v / ln["qty"] if ln["qty"] else 1.0)
            amount = round(ln["amount"] * frac, 2)
        elif l["basis"] == "AMOUNT":
            need, qty, frac, amount = v, 0.0, (v / ln["amount"] if ln["amount"] else 1.0), round(v, 2)
        else:   # SINGLE: the one and only invoice, whole line
            need, qty, frac, amount = 1.0, 1.0, 1.0, round(ln["amount"], 2)
            if pid in single_used:
                errors.append(f"{label}: a single-invoice PO/PR line can be used only once")
                continue
            single_used.add(pid)
        tol = 1e-9 if l["basis"] in ("QTY", "SINGLE") else 0.005
        if need > local[pid] + tol:
            errors.append(f"{label}: only {max(local[pid], 0):g} left to bill, you entered {need:g}")
            continue
        local[pid] -= need
        cover[i] += frac
        is_cited = _ref_norm(l["po_reference"]) in cited or ln["po"] == _ref_norm(l["po_reference"])
        sole = open_per_po.get(l["po_record_id"], 0) == 1
        score = line_score(ln, l, is_cited, sole and is_cited)
        vs = round(text_sim(p.get("vendor_name"), l["vendor"]) or 0.0, 3)
        warn = []
        if score < LINK_THRESHOLD:
            warn.append(f"automatic score {score:.2f} is below {LINK_THRESHOLD}")
        if vs < VENDOR_MIN:
            warn.append(f"vendor name looks different (similarity {vs:.2f})")
        out.append(dict(inv_line=i, po_line_id=pid, po_reference=l["po_reference"], po_item=l["description"], basis=l["basis"],
                        applied=need if l["basis"] != "SINGLE" else "whole line", qty=qty, amount=amount, score=score,
                        vendor_similarity=vs, warning="; ".join(warn)))
    if not rows:
        errors.append("Nothing selected")
    for i, cv in cover.items():
        if abs(cv - 1.0) > 0.0005:
            errors.append(f"Invoice line {i} is {cv * 100:.1f}% allocated; it must be 100%")
    result = dict(ok=not errors, errors=errors, rows=out, below_threshold=any(r["warning"] for r in out))
    if dry_run:
        return result
    if errors:
        raise ValueError("; ".join(errors))
    if len((reason or "").strip()) < 5:
        raise ValueError("A typed reason (at least 5 characters) is required for a manual match")
    now = _now()
    for r in out:
        con.execute("INSERT INTO allocations(doc_record_id, case_id, kind, po_line_id, qty, amount, score, created_at, manual, reason) "
                    "VALUES(?,?,?,?,?,?,?,?,1,?)", (case["inv_record_id"], case_id, INV, r["po_line_id"], r["qty"], r["amount"], r["score"], now, reason.strip()))
        con.execute("UPDATE po_lines SET billed_qty=billed_qty+?, billed_amount=billed_amount+? WHERE line_id=?",
                    (r["qty"], r["amount"], r["po_line_id"]))
    score = round(sum(r["score"] for r in out) / len(out), 4)
    con.execute("UPDATE cases SET reason='', link_score=?, hold=0, updated_at=? WHERE case_id=?", (score, now, case_id))
    con.execute("UPDATE documents SET outcome='ALLOCATED' WHERE record_id=?", (case["inv_record_id"],))
    status = refresh_case(con, case_id)
    warns = [f"line {r['inv_line']}: {r['warning']}" for r in out if r["warning"]]
    _log(con, case["inv_record_id"], case_id, "MANUAL_MATCH",
         f"{len(out)} allocation(s) to {', '.join(sorted({r['po_reference'] for r in out}))}; status={status}; reason: {reason.strip()}"
         + (f"; warnings: {' | '.join(warns)}" if warns else ""))
    con.commit()
    return dict(ok=True, status=status, score=score, warnings=warns)

def po_detail(con, record_id: str) -> dict:
    """Everything the PO preview page shows: the document, its lines with balances, and the invoices charged to it."""
    doc = get_doc(con, record_id)
    modes = {"QTY": "Goods (by quantity)", "AMOUNT": "Services (by amount)", "SINGLE": "Services (single invoice)"}
    lines = []
    for l in con.execute("SELECT * FROM po_lines WHERE po_record_id=? ORDER BY line_no", (record_id,)):
        l = dict(l)
        qty = l["basis"] in ("QTY", "SINGLE")
        lines.append(dict(line=l["line_no"], item=l["description"], basis=modes[l["basis"]],
                          ordered=l["ordered_qty"] if qty else l["ordered_amount"],
                          billed=l["billed_qty"] if qty else l["billed_amount"],
                          received=l["received_qty"] if qty else l["received_amount"],
                          left_to_bill=round(_open(l), 4)))
    invs = []
    for a in con.execute("SELECT a.*, c.status, c.inv_record_id FROM allocations a JOIN po_lines l ON l.line_id=a.po_line_id "
                         "JOIN cases c ON c.case_id=a.case_id WHERE l.po_record_id=? ORDER BY a.alloc_id", (record_id,)):
        ip = get_doc(con, a["inv_record_id"])["payload"]
        invs.append(dict(case_id=a["case_id"], invoice=ip.get("invoice_reference"), status=a["status"], qty=a["qty"],
                         amount=a["amount"], matched="manual" if a["manual"] else "automatic"))
    return dict(doc=doc, lines=lines, invoices=invs)

# =====================================================================================
# 14D. CALL-OFFS (PO/PRs that cite no amount: their own case, manual matching to several invoices)
# =====================================================================================
def calloff_rows(con, statuses) -> list[dict]:
    rows = []
    if not statuses:
        return rows
    q = f"SELECT * FROM calloffs WHERE status IN ({','.join('?' * len(statuses))}) ORDER BY case_id"
    for c in con.execute(q, list(statuses)).fetchall():
        po = get_doc(con, c["po_record_id"])["payload"]
        invs = [get_doc(con, r["inv_record_id"])["payload"] for r in
                con.execute("SELECT inv_record_id FROM calloff_invoices WHERE case_id=? ORDER BY id", (c["case_id"],))]
        total = round(sum(float(doc_total(p) or 0) for p in invs), 2) if invs else None
        rows.append(dict(case_id=c["case_id"], vendor=po.get("vendor_name"),
                         invoice_reference=", ".join(str(p.get("invoice_reference")) for p in invs) or "-",
                         po_references=po.get("po_reference"), amount=total, link_score=None, status=c["status"],
                         reason=c["reason"] or "", updated_at=c["updated_at"], kind="Call-Off", type=po.get("category"),
                         matched=len(invs), date=po.get("date")))
    return rows

def calloff_detail(con, case_id: str) -> dict:
    c = con.execute("SELECT * FROM calloffs WHERE case_id=?", (case_id,)).fetchone()
    if not c:
        raise ValueError("Unknown call-off case")
    invs = [dict(case_id=r["inv_case_id"], doc=get_doc(con, r["inv_record_id"])) for r in
            con.execute("SELECT * FROM calloff_invoices WHERE case_id=? ORDER BY id", (case_id,))]
    return dict(case=dict(c), po=get_doc(con, c["po_record_id"]), invoices=invs)

def calloff_open_invoices(con, case_id: str) -> list[dict]:
    """Invoices still waiting in Stage 1, most similar vendor first, for the dropdown."""
    c = con.execute("SELECT * FROM calloffs WHERE case_id=?", (case_id,)).fetchone()
    vendor = get_doc(con, c["po_record_id"])["payload"].get("vendor_name")
    out = []
    for r in con.execute("SELECT * FROM cases WHERE status=? AND calloff_id IS NULL ORDER BY case_id", (S_OPEN_PO,)):
        p = get_doc(con, r["inv_record_id"])["payload"]
        vs = round(text_sim(vendor, p.get("vendor_name")) or 0.0, 3)
        amt = doc_total(p)
        out.append(dict(case_id=r["case_id"], vendor_sim=vs, vendor=p.get("vendor_name"), invoice_reference=p.get("invoice_reference"),
                        label=f"{r['case_id']} | {p.get('invoice_reference')} | {p.get('vendor_name')} | "
                              f"{amt:,.2f} | {p.get('date')}" if amt is not None else f"{r['case_id']} | {p.get('invoice_reference')} | {p.get('vendor_name')}"))
    out.sort(key=lambda x: (-x["vendor_sim"], x["case_id"]))
    return out

def confirm_calloff(con, case_id: str, inv_case_ids: list, reason: str) -> dict:
    """Match the chosen invoices to this call-off PO/PR. The case moves to Stage 2 (Pending GR/SES) until receipts cover them."""
    c = con.execute("SELECT * FROM calloffs WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] != S_CALLOFF:
        raise ValueError("This call-off case is not open")
    ids = [i for i in inv_case_ids if i]
    if not ids:
        raise ValueError("Select at least one invoice")
    if len(set(ids)) != len(ids):
        raise ValueError("The same invoice is selected more than once")
    if len((reason or "").strip()) < 5:
        raise ValueError("A typed reason (at least 5 characters) is required for a manual match")
    po = get_doc(con, c["po_record_id"])["payload"]
    rows, warns = [], []
    for i in ids:
        r = con.execute("SELECT * FROM cases WHERE case_id=?", (i,)).fetchone()
        if not r or r["status"] != S_OPEN_PO or r["calloff_id"]:
            raise ValueError(f"{i} is no longer an open invoice")
        rows.append(r)
        p = get_doc(con, r["inv_record_id"])["payload"]
        vs = text_sim(po.get("vendor_name"), p.get("vendor_name")) or 0.0
        if vs < VENDOR_MIN:
            warns.append(f"{i}: vendor name looks different (similarity {vs:.2f})")
    now = _now()
    for r in rows:
        con.execute("INSERT INTO calloff_invoices(case_id, inv_case_id, inv_record_id, added_at) VALUES(?,?,?,?)",
                    (case_id, r["case_id"], r["inv_record_id"], now))
        con.execute("UPDATE cases SET status=?, calloff_id=?, reason=?, hold=0, updated_at=? WHERE case_id=?",
                    (S_CALLED, case_id, f"Matched to call-off {case_id}", now, r["case_id"]))
        con.execute("UPDATE documents SET outcome='CALLOFF_MATCHED', case_id=? WHERE record_id=?", (r["case_id"], r["inv_record_id"]))
    con.execute("UPDATE calloffs SET status=?, reason=?, updated_at=? WHERE case_id=?", (S_GR, reason.strip(), now, case_id))
    _log(con, c["po_record_id"], case_id, "CALLOFF_CONFIRMED",
         f"{po.get('po_reference')} matched to {len(rows)} invoice(s): {', '.join(r['case_id'] for r in rows)}; reason: {reason.strip()}"
         + (f"; warnings: {' | '.join(warns)}" if warns else ""))
    con.commit()
    return dict(case_id=case_id, status=S_GR, invoices=[r["case_id"] for r in rows], warnings=warns)

def _calloff_receipts(con, case_id: str) -> dict:
    """GR/SES receipts raised for this call-off case: (invoice case, invoice line) -> [qty, amount]."""
    out = {}
    for r in con.execute("SELECT payload FROM documents WHERE doc_type=? AND case_id=? AND outcome='GR_SUBMITTED'", (GR, case_id)):
        for x in json.loads(r["payload"]).get("lines", []):
            k = (x["inv_case_id"], int(x["inv_line"]))
            q, a_ = out.get(k, [0.0, 0.0])
            out[k] = [q + float(x.get("qty") or 0), a_ + float(x.get("amount") or 0)]
    return out

def calloff_lines(con, case_id: str) -> list[dict]:
    """Lines of every invoice matched to a call-off case, with what is billed and what has been received so far.
    Goods invoice lines are received by quantity, service lines by amount."""
    got, out = _calloff_receipts(con, case_id), []
    for r in con.execute("SELECT * FROM calloff_invoices WHERE case_id=? ORDER BY id", (case_id,)).fetchall():
        p = get_doc(con, r["inv_record_id"])["payload"]
        for i, ln in enumerate(_invoice_lines(p), 1):
            qty = bool(p.get("items"))
            billed = ln["qty"] if qty else ln["amount"]
            received = got.get((r["inv_case_id"], i), [0.0, 0.0])[0 if qty else 1]
            out.append(dict(inv_case_id=r["inv_case_id"], invoice_reference=p.get("invoice_reference"), inv_line=i,
                            description=ln["description"], basis="QTY" if qty else "AMOUNT", unit_price=ln["unit_price"],
                            line_amount=ln["amount"], billed=billed, received=received, needed=round(max(0.0, billed - received), 4)))
    return out

def refresh_calloff(con, case_id: str) -> str:
    """Pending GR/SES until receipts cover every line of every matched invoice, then Complete."""
    c = con.execute("SELECT * FROM calloffs WHERE case_id=?", (case_id,)).fetchone()
    if c["status"] in (S_ARCHIVED, S_CALLOFF):
        return c["status"]
    lines = calloff_lines(con, case_id)
    status = S_DONE if lines and all(l["needed"] <= (1e-9 if l["basis"] == "QTY" else 0.005) for l in lines) else S_GR
    con.execute("UPDATE calloffs SET status=?, updated_at=? WHERE case_id=?", (status, _now(), case_id))
    return status

def submit_calloff_gr(con, case_id: str, gr: dict) -> dict:
    """One GR/SES batch for a call-off case. gr = {service_date_from, service_date_to, prepared_by,
    lines:[{inv_case_id, inv_line, receive}]}; receive = quantity (goods lines) or amount (service lines)."""
    c = con.execute("SELECT * FROM calloffs WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] != S_GR:
        raise ValueError("This call-off case is not pending GR/SES")
    lines = {(l["inv_case_id"], l["inv_line"]): l for l in calloff_lines(con, case_id)}
    wanted = [(x["inv_case_id"], int(x["inv_line"]), float(x.get("receive") or 0)) for x in gr.get("lines", []) if float(x.get("receive") or 0) > 0]
    if not wanted:
        raise ValueError("Enter a received quantity or amount for at least one line")
    po = get_doc(con, c["po_record_id"])["payload"]
    if _norm(gr.get("prepared_by")) != _norm(po.get("prepared_by")):
        raise ValueError(f"Prepared by must match the PO/PR: {po.get('prepared_by')}")
    d1, d2 = parse_date(gr.get("service_date_from")), parse_date(gr.get("service_date_to"))
    if not d1 or not d2:
        raise ValueError("Performance / service dates are required")
    if d2 < d1:
        raise ValueError("Performance end date is before start date")
    out_lines, total = [], 0.0
    for icase, iline, recv in wanted:
        l = lines.get((icase, iline))
        if not l:
            raise ValueError("Line does not belong to this call-off case")
        if recv > l["needed"] + (1e-9 if l["basis"] == "QTY" else 0.005):
            raise ValueError(f"{l['invoice_reference']} '{(l['description'] or '')[:30]}': receiving {recv:g} exceeds the {l['needed']:g} still needed")
        amount = round(l["line_amount"] * recv / l["billed"], 2) if l["basis"] == "QTY" and l["billed"] else round(recv, 2)
        total += amount
        out_lines.append(dict(inv_case_id=icase, invoice_reference=l["invoice_reference"], inv_line=iline, description=l["description"],
                              basis=l["basis"], qty=recv if l["basis"] == "QTY" else None, amount=amount))
    category = po.get("category") or "Services"
    number = next_ses_number(con, category)
    payload = dict(doc_type=GR, record_id=number, ses_gr_number=number, case_id=case_id, vendor_name=po.get("vendor_name"), category=category,
                   po_references=[po.get("po_reference")], service_date_from=gr["service_date_from"], service_date_to=gr["service_date_to"],
                   prepared_by=gr["prepared_by"], lines=out_lines, total=round(total, 2))
    con.execute("INSERT INTO documents VALUES(?,?,?,?,?,?,?,?,?,?)",
                (number, GR, json.dumps(payload), None, _now(), "GR_SUBMITTED", case_id, None, None, "[]"))
    status = refresh_calloff(con, case_id)
    _log(con, number, case_id, "GR_SUBMITTED", f"{number}: {len(out_lines)} line(s), total {payload['total']}; call-off status={status}")
    con.commit()
    return dict(record_id=number, case_id=case_id, status=status, ses_gr_number=number)

def _build_calloff_package(con, case_id: str):
    d = calloff_detail(con, case_id)
    buf, files = io.BytesIO(), []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for label, doc in [("po_pr", d["po"])] + [("invoice", i["doc"]) for i in d["invoices"]]:
            p = Path(doc["file_path"]) if doc and doc.get("file_path") else None
            if p and p.exists():
                name = f"{case_id}/{label}_{_orig_name(p)}"
                z.write(p, name)
                files.append(name)
        for g in case_gr_docs(con, case_id):
            name = f"{case_id}/gr_ses_{g['payload']['ses_gr_number']}.json"
            z.writestr(name, json.dumps(g["payload"], indent=2))
            files.append(name)
        po = d["po"]["payload"]
        summary = dict(case_id=case_id, kind="call-off", status=d["case"]["status"], vendor=po.get("vendor_name"),
                       po_reference=po.get("po_reference"), reason=d["case"]["reason"],
                       invoices=[dict(case_id=i["case_id"], invoice_reference=i["doc"]["payload"].get("invoice_reference"),
                                      total=doc_total(i["doc"]["payload"])) for i in d["invoices"]],
                       gr_ses=[g["payload"]["ses_gr_number"] for g in case_gr_docs(con, case_id)],
                       files=files, packaged_at=_now())
        z.writestr(f"{case_id}/case_summary.json", json.dumps(summary, indent=2))
    return buf.getvalue(), f"{case_id}.zip"

def _post_calloff(con, case_id: str) -> dict:
    c = con.execute("SELECT * FROM calloffs WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] != S_DONE:
        raise ValueError("Only Complete call-off cases can be posted")
    data, name = _build_calloff_package(con, case_id)
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    path = ARCHIVE_DIR / name
    path.write_bytes(data)
    now = _now()
    con.execute("INSERT OR REPLACE INTO archive VALUES(?,?,?)", (case_id, now, str(path)))
    con.execute("UPDATE calloffs SET status=?, updated_at=? WHERE case_id=?", (S_ARCHIVED, now, case_id))
    con.execute("UPDATE cases SET status=?, updated_at=? WHERE calloff_id=?", (S_ARCHIVED, now, case_id))
    _log(con, None, case_id, "POSTED", f"call-off package saved to Complete Archive ({name})")
    con.commit()
    return dict(case_id=case_id, posted_at=now, package=str(path))

# =====================================================================================
# 14B. POST AND COMPLETE ARCHIVE
# =====================================================================================
def _orig_name(path) -> str:
    return re.sub(r"^[0-9a-f]{10}_", "", Path(path).name)

def build_case_package(con, case_id: str):
    """Zip: invoice, every matched PO/PR (original format), every GR/SES batch as JSON, and a summary."""
    if str(case_id).startswith("CALL-"):
        return _build_calloff_package(con, case_id)
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not c:
        raise ValueError("Unknown case")
    inv, pos, grs = get_doc(con, c["inv_record_id"]), case_po_docs(con, case_id), case_gr_docs(con, case_id)
    buf, files = io.BytesIO(), []
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for label, doc in [("invoice", inv)] + [("po_pr", d) for d in pos]:
            p = Path(doc["file_path"]) if doc and doc.get("file_path") else None
            if p and p.exists():
                name = f"{case_id}/{label}_{_orig_name(p)}"
                z.write(p, name)
                files.append(name)
        for g in grs:
            name = f"{case_id}/gr_ses_{g['payload']['ses_gr_number']}.json"
            z.writestr(name, json.dumps(g["payload"], indent=2))
            files.append(name)
        z.writestr(f"{case_id}/allocations.json", json.dumps(
            [{k: l[k] for k in ("po_reference", "description", "basis", "billed_this", "billed_cum", "ordered", "received", "score")}
             for l in case_lines(con, case_id)], indent=2))
        summary = dict(case_id=case_id, status=c["status"], link_score=c["link_score"], vendor=inv["payload"].get("vendor_name"),
                       invoice_reference=inv["payload"].get("invoice_reference"), invoice_total=doc_total(inv["payload"]),
                       po_references=[d["payload"].get("po_reference") for d in pos],
                       gr_ses=[g["payload"]["ses_gr_number"] for g in grs], files=files, packaged_at=_now())
        z.writestr(f"{case_id}/case_summary.json", json.dumps(summary, indent=2))
    return buf.getvalue(), f"{case_id}.zip"

def post_case(con, case_id: str) -> dict:
    if str(case_id).startswith("CALL-"):
        return _post_calloff(con, case_id)
    c = con.execute("SELECT * FROM cases WHERE case_id=?", (case_id,)).fetchone()
    if not c or c["status"] != S_DONE:
        raise ValueError("Only Complete cases can be posted")
    data, name = build_case_package(con, case_id)
    ARCHIVE_DIR.mkdir(parents=True, exist_ok=True)
    path = ARCHIVE_DIR / name
    path.write_bytes(data)
    now = _now()
    con.execute("INSERT OR REPLACE INTO archive VALUES(?,?,?)", (case_id, now, str(path)))
    con.execute("UPDATE cases SET status=?, updated_at=? WHERE case_id=?", (S_ARCHIVED, now, case_id))
    _log(con, None, case_id, "POSTED", f"package saved to Complete Archive ({name})")
    con.commit()
    return dict(case_id=case_id, posted_at=now, package=str(path))

def read_archive_package(con, case_id: str):
    a = con.execute("SELECT package_path FROM archive WHERE case_id=?", (case_id,)).fetchone()
    p = Path(a["package_path"]) if a else None
    return p.read_bytes() if p and p.exists() else None

# =====================================================================================
# 15. READ HELPERS FOR THE UI
# =====================================================================================
def case_rows(con, statuses) -> list[dict]:
    """Invoice cases for the given statuses, plus Call-Off cases whose status is in the list."""
    q = f"SELECT * FROM cases WHERE status IN ({','.join('?' * len(statuses))}) AND calloff_id IS NULL ORDER BY case_id"
    rows = []
    for c in con.execute(q, list(statuses)):
        p = get_doc(con, c["inv_record_id"])["payload"]
        pos = sorted({r[0] for r in con.execute("SELECT DISTINCT l.po_reference FROM allocations a JOIN po_lines l ON l.line_id=a.po_line_id "
                                                "WHERE a.case_id=?", (c["case_id"],))}) or list(_cited(p).values())
        rows.append(dict(case_id=c["case_id"], vendor=p.get("vendor_name"), invoice_reference=p.get("invoice_reference"),
                         po_references=", ".join(pos), amount=doc_total(p), link_score=c["link_score"], status=c["status"],
                         reason=c["reason"] or "", updated_at=c["updated_at"]))
    return rows + calloff_rows(con, [s for s in statuses if s in (S_CALLOFF, S_GR, S_DONE, S_ARCHIVED)])

def archive_rows(con) -> list[dict]:
    rows = case_rows(con, [S_ARCHIVED])
    for r in rows:
        a = con.execute("SELECT * FROM archive WHERE case_id=?", (r["case_id"],)).fetchone()
        r["posted_at"] = a["posted_at"] if a else None
    return rows

def po_rows(con) -> list[dict]:
    """PO/PR positions: what is ordered, billed and received, and what can still be billed."""
    rows = []
    modes = {"QTY": "Goods (by quantity)", "AMOUNT": "Services (by amount)", "SINGLE": "Services (single invoice)"}
    for r in con.execute("SELECT record_id FROM documents WHERE doc_type=? ORDER BY received_at, record_id", (PO,)):
        d = get_doc(con, r[0])
        ls = [dict(x) for x in con.execute("SELECT * FROM po_lines WHERE po_record_id=? ORDER BY line_no", (r[0],))]
        if not ls:   # Call-Off PO/PR: handled in the Call-Offs tab
            continue
        single = all(l["basis"] == "SINGLE" for l in ls)
        ordered = sum(l["ordered_amount"] for l in ls)
        open_bill = sum(max(0.0, _open(l)) * (l["unit_price"] if l["basis"] == "QTY" else 1.0) for l in ls if l["basis"] != "SINGLE")
        any_open = any(_open(l) > (1e-9 if l["basis"] in ("QTY", "SINGLE") else 0.005) for l in ls)
        billed = sum(l["billed_qty"] * l["unit_price"] if l["basis"] == "QTY" else l["billed_amount"] for l in ls)
        billed_any = any(l["billed_qty"] > 0 or l["billed_amount"] > 0 for l in ls)
        received = sum(l["received_qty"] * l["unit_price"] if l["basis"] == "QTY" else l["received_amount"] for l in ls)
        status = S_OPEN_INV if not billed_any else (S_PART if any_open else S_BILLED)
        if not any_open and not billed_any:
            status = S_BILLED
        p = d["payload"]
        rows.append(dict(record_id=r[0], po_reference=p.get("po_reference"), vendor=p.get("vendor_name"), mode=modes[ls[0]["basis"]],
                         status=status, ordered=None if single else round(ordered, 2), billed=round(billed, 2), received=round(received, 2),
                         open_to_bill=None if single else round(open_bill, 2), date=p.get("date")))
    return rows

def counts(con) -> dict:
    """Pending entries for the tab bubbles."""
    n = lambda q, *a: con.execute(q, a).fetchone()[0]
    open_pos = sum(1 for r in po_rows(con) if r["status"] != S_BILLED)
    return dict(review=n("SELECT COUNT(*) FROM intake WHERE route=?", R_REVIEW),
                s1=n("SELECT COUNT(*) FROM cases WHERE status=?", S_OPEN_PO) + open_pos,
                s2=n("SELECT COUNT(*) FROM cases WHERE status=?", S_GR) + n("SELECT COUNT(*) FROM calloffs WHERE status=?", S_GR),
                s3=n("SELECT COUNT(*) FROM cases WHERE status=?", S_DONE) + n("SELECT COUNT(*) FROM calloffs WHERE status=?", S_DONE),
                calloff=n("SELECT COUNT(*) FROM calloffs WHERE status=?", S_CALLOFF))

def list_documents(con) -> list[dict]:
    """One row per record across Module 1 (channel, status, route) and Module 2 (outcome, case)."""
    rows = []
    q = """SELECT i.record_id, i.channel, i.source_file, i.doc_type AS m1_type, i.m1_status, i.route, i.missing AS m1_missing,
                  i.received_at, d.doc_type, d.outcome AS m2_outcome, d.case_id, d.missing AS m2_missing
           FROM intake i LEFT JOIN documents d ON d.record_id = i.record_id ORDER BY i.received_at"""
    for r in con.execute(q):
        miss = json.loads(r["m2_missing"] or "[]") or json.loads(r["m1_missing"] or "[]")
        rows.append(dict(record_id=r["record_id"], doc_type=r["doc_type"] or r["m1_type"], channel=r["channel"],
                         source_file=r["source_file"], module1_status=r["m1_status"], route=r["route"],
                         module2_outcome=r["m2_outcome"], case_id=r["case_id"], missing_fields=", ".join(miss),
                         received_at=r["received_at"]))
    for r in con.execute("SELECT * FROM documents WHERE doc_type=?", (GR,)):
        rows.append(dict(record_id=r["record_id"], doc_type=GR, channel="GR/SES FORM", source_file=None, module1_status=None,
                         route=None, module2_outcome=r["outcome"], case_id=r["case_id"], missing_fields="", received_at=r["received_at"]))
    return rows

def list_log(con, stages=None) -> list[dict]:
    q, args = "SELECT * FROM log", []
    if stages:
        q += f" WHERE stage IN ({','.join('?' * len(stages))})"
        args = list(stages)
    return [dict(r) for r in con.execute(q + " ORDER BY ts DESC, rowid DESC", args)]
